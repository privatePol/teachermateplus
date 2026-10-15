from datetime import date
from decimal import Decimal
import json

from django.conf import settings
from django.test import SimpleTestCase, TestCase
from django.db import connection
from django.test.utils import CaptureQueriesContext
from unittest.mock import patch
from urllib.parse import parse_qs

from apps.admin_portal.grade_distribution import GradeDistributionMonitorService
from apps.admin_portal.grade_distribution_metrics import FormulaPrivacyPlan, PrivacyFamily, distribution, metrics
from apps.admin_portal.grade_distribution_selection import logical_period
from apps.grading.services import FacultyGradingService
from django.urls import reverse
from django.utils import timezone

from apps.accounts.models import User
from apps.academics.models import AcademicYear, Course, CourseOffering, FacultyAssignment, Section, Term
from apps.core.services.settings import SystemSettingService
from apps.enrollment.models import Enrollment
from apps.grading.models import (
    GradeActivity,
    GradingTemplate,
    GradingTemplateComponent,
    GradingTemplatePeriod,
    StudentActivityScore,
    StudentPeriodGrade,
    TenantGradingProfile,
    GradingTemplateSubcomponent,
    GradingTemplateDetail,
)
from apps.rbac.models import Permission, Role, RolePermission, UserRole, UserPermission
from apps.students.models import Student
from apps.tenants.models import Campus, Department, Program, Tenant


class GradeDistributionMonitorTests(TestCase):
    def setUp(self):
        fixture_clock = patch("apps.admin_portal.grade_distribution.timezone.localdate", return_value=date(2025, 7, 1))
        fixture_clock.start()
        self.addCleanup(fixture_clock.stop)
        self.tenant = Tenant.objects.create(code="NCBA", name="NCBA")
        self.campus = Campus.objects.create(tenant=self.tenant, code="FV", name="Fairview")
        self.department = Department.objects.create(
            tenant=self.tenant,
            campus=self.campus,
            code="BSA",
            name="Accountancy",
        )
        self.program = Program.objects.create(
            tenant=self.tenant,
            campus=self.campus,
            department=self.department,
            code="BSA",
            name="BS Accountancy",
        )
        self.academic_year = AcademicYear.objects.create(
            tenant=self.tenant,
            code="2025-2026",
            name="AY 2025-2026",
            start_date=date(2025, 6, 1),
            end_date=date(2026, 5, 31),
        )
        self.term = Term.objects.create(
            tenant=self.tenant,
            academic_year=self.academic_year,
            code="1ST",
            name="First Semester",
            sequence_no=1,
            start_date=date(2025, 6, 1),
            end_date=date(2025, 10, 31),
        )
        self.template = GradingTemplate.objects.create(
            tenant=self.tenant,
            code="GENED",
            name="General Education",
            is_published=True,
            is_active=True,
            passing_grade_threshold=Decimal("75"),
        )
        self.period = GradingTemplatePeriod.objects.create(
            template=self.template,
            code="PRELIM",
            name="Prelim",
            sequence_no=1,
        )

        self.course = Course.objects.create(
            tenant=self.tenant,
            campus=self.campus,
            department=self.department,
            code="ACC101",
            title="Accounting 101",
        )
        self.section = Section.objects.create(
            tenant=self.tenant,
            campus=self.campus,
            department=self.department,
            program=self.program,
            code="BSA1A",
            name="BSA 1A",
        )
        self.offering = CourseOffering.objects.create(
            tenant=self.tenant,
            campus=self.campus,
            department=self.department,
            program=self.program,
            academic_year=self.academic_year,
            term=self.term,
            course=self.course,
            section=self.section,
        )

        self.admin_user = self._user("dean", "Dean", "User", self.tenant, self.campus, self.department)
        self.faculty_user = self._user("faculty", "Faculty", "One", self.tenant, self.campus, self.department)
        self._assign_permissions(self.admin_user, include_monitor=True)
        self._assign_faculty_role(self.faculty_user)
        FacultyAssignment.objects.create(
            tenant=self.tenant,
            campus=self.campus,
            offering=self.offering,
            faculty_user=self.faculty_user,
            response_status=FacultyAssignment.ResponseStatus.ACCEPTED,
            accepted_at=timezone.now(),
            is_primary=True,
        )

        self._seed_grades(self.offering, [95, 94, 93, 92, 91, 90, 96, 97, 80, 76])
        self.url = reverse("admin_portal:grade_distribution_monitor")

    def search(self, **kwargs):
        return self.client.get(self.url, {"apply": "1", "academic_year_id": self.academic_year.id, "term_id": self.term.id, "period": "PRELIM", **kwargs})

    def detail(self, response, row=None):
        row = row or response.context["rows"][0]
        url = reverse("admin_portal:grade_distribution_monitor_detail", args=[row["offering_id"], row["physical_period_id"]])
        return self.client.get(url + "?" + row["detail_query"])

    def test_monitor_requires_specific_permission(self):
        user = self._user("no_monitor", "No", "Monitor", self.tenant, self.campus, self.department)
        self._assign_permissions(user, include_monitor=False)
        self.client.force_login(user)
        self.assertEqual(self.client.get(self.url).status_code, 403)
        self.assertEqual(self.client.get(reverse("admin_portal:grade_distribution_monitor_detail", args=[self.offering.id, self.period.id])).status_code, 403)
        self.assertEqual(self.client.get(self.url, {"export": "csv"}).status_code, 403)

    def test_initial_load_and_unsubmitted_filters_do_not_query_grade_results(self):
        self.client.force_login(self.admin_user)
        for params in ({}, {"academic_year_id": self.academic_year.id, "term_id": self.term.id, "period": "PRELIM"}, {"options": "1"}):
            with CaptureQueriesContext(connection) as queries:
                response = self.client.get(self.url, params)
            self.assertEqual(response.status_code, 200)
            self.assertFalse(any('student_period_grades' in q["sql"] or 'student_activity_scores' in q["sql"] or 'grade_activities' in q["sql"] for q in queries))
        self.assertContains(self.client.get(self.url), "Select criteria and click Apply Filters")

    def test_required_and_malformed_selections_fail_before_grade_retrieval(self):
        self.client.force_login(self.admin_user)
        for params in ({"apply": "1"}, {"apply": "1", "term_id": "bad"}, {"apply": "1", "academic_year_id": "999999", "term_id": self.term.id, "period": "FINAL"}):
            with CaptureQueriesContext(connection) as queries:
                response = self.client.get(self.url, params)
            self.assertTrue(response.context["errors"])
            self.assertEqual(response.context["rows"], [])
            self.assertFalse(any('student_period_grades' in q["sql"] for q in queries))

    def test_term_must_belong_to_year(self):
        other_year = AcademicYear.objects.create(tenant=self.tenant, code="OLD", name="Old", start_date=date(2024, 1, 1), end_date=date(2024, 12, 31))
        self.client.force_login(self.admin_user)
        self.assertTrue(self.search(academic_year_id=other_year.id).context["errors"])

    def test_summary_is_aggregate_and_average_uses_saved_values(self):
        self.client.force_login(self.admin_user)
        response = self.search()
        self.assertEqual(response.status_code, 200)
        row = response.context["rows"][0]
        self.assertEqual(row["average"], Decimal("90.40"))
        self.assertEqual(row["graded_count"], 10)
        self.assertEqual(row["status"], "Unfinalized")
        self.assertEqual(sum(b["count"] for b in row["distribution"]), 10)
        self.assertContains(response, "Average Grade")
        self.assertNotContains(response, "Review Flags")
        self.assertNotContains(response, "Masked Student")
        self.assertNotIn("grade_details", row)

    def test_high_grade_html_json_xlsx_parity_and_unrounded_denominators(self):
        grades = list(StudentPeriodGrade.objects.filter(offering=self.offering).order_by("id"))
        StudentPeriodGrade.objects.filter(offering=self.offering).update(exam_grade=None, class_standing_grade=None)
        for grade, value in zip(grades, (100, 99, 98, 97, 96, 95, Decimal("99.75"), 90, 89, None)):
            StudentPeriodGrade.objects.filter(pk=grade.pk).update(period_grade=value)
        self.client.force_login(self.admin_user)
        response = self.search()
        row = response.context["rows"][0]
        data = self.detail(response).json()
        export = self._privacy_xlsx_row(self.client.get(self.url + "?" + response.context["export_query"]))
        self.assertEqual([item["count"] for item in row["high_grade_breakdown"]], [1, 1, 1, 1, 1, 1, 6, 8])
        self.assertContains(response, "gdm-metric gdm-exact-100")
        self.assertContains(response, "11.11%")
        for local, public in zip(row["high_grade_breakdown"], data["high_grade_breakdown"]):
            label = local["grade"]
            self.assertEqual(public["count"], local["count"])
            self.assertEqual(public["total_graded_students"], 9)
            self.assertEqual(Decimal(public["percentage"]), local["percentage"])
        self.assertTrue(all(export[f"Exactly {grade}"] == "" for grade in range(100, 94, -1)))
        self.assertNotIn("Actual examination percentage", json.dumps(data))
        self.assertEqual(data["missing_count"], 1)

    def test_gradebook_activity_equivalence_requires_full_maps_not_matching_averages(self):
        exam = GradingTemplateComponent.objects.create(template_period=self.period, code="EXAM", name="Exam", weight_percentage=100, is_exam_component=True)
        activity = GradeActivity.objects.create(tenant=self.tenant, campus=self.campus, offering=self.offering,
            template_period=self.period, template_component=exam, title="Midterm Exam", total_score=100)
        for grade in StudentPeriodGrade.objects.filter(offering=self.offering):
            StudentActivityScore.objects.create(activity=activity, student_id=grade.student_id, raw_score=50, computed_score=100)
        StudentPeriodGrade.objects.filter(offering=self.offering).update(period_grade=100, exam_grade=100, class_standing_grade=None)
        self.client.force_login(self.admin_user)
        data = self.detail(self.search()).json()
        self.assertEqual([m["label"] for m in data["measures"]], ["Examination Grades", "Class Standing Grades", "Midterm Exam"])
        configured = data["measures"][-1]
        self.assertEqual(configured["equivalent_to"], "Period Grades")
        self.assertEqual(configured["high_grade_breakdown"][0]["count"], 10)
        scores = list(StudentActivityScore.objects.filter(activity=activity).order_by("id"))
        for index, score in enumerate(scores):
            StudentActivityScore.objects.filter(pk=score.pk).update(computed_score=99 if index < 5 else 101)
        configured = self.detail(self.search()).json()["measures"][-1]
        self.assertEqual(Decimal(configured["average"]), 100)
        self.assertEqual(configured["equivalent_to"], "")
        self.assertEqual(configured["high_grade_breakdown"][0]["count"], 0)

    def test_modal_template_has_accessible_responsive_collapsible_and_loading_styles(self):
        self.client.force_login(self.admin_user)
        response = self.client.get(self.url)
        for phrase in ("prefers-reduced-motion", ":focus-visible", "#E6F3E9", "#FAFDFB", "#D0E5D6",
                       "@keyframes gdm-progress", "modal-fullscreen-sm-down", 'aria-labelledby="gdm-modal-title"',
                       'id="gdm-filter-status" role="status" aria-live="polite"'):
            self.assertContains(response, phrase)
        self.assertNotContains(response, 'role="progressbar"')  # Initial page remains idle.

    def test_all_four_logical_periods_and_exact_options(self):
        students = list(StudentPeriodGrade.objects.filter(offering=self.offering).values_list("student_id", flat=True))
        for key, name, seq in (("MIDTERM", "Midterm", 2), ("PRE-FINAL", "Pre-Final", 3), ("FINAL_EXAM", "Final Exam", 4)):
            p = GradingTemplatePeriod.objects.create(template=self.template, code=key, name=name, sequence_no=seq)
            StudentPeriodGrade.objects.create(tenant=self.tenant, campus=self.campus, offering=self.offering, template_period=p, student_id=students[0], period_grade=Decimal("100"))
        self.client.force_login(self.admin_user)
        options = self.client.get(self.url).context["filter_options"]["periods"]
        self.assertEqual([p["label"] for p in options], ["Prelim", "Midterm", "Pre-Final", "Final"])
        for key in ("PRELIM", "MIDTERM", "PREFINAL", "FINAL"):
            self.assertEqual(len(self.search(period=key).context["rows"]), 1)

    def test_ambiguous_period_not_guessed_and_distinct_physical_records_not_merged(self):
        unknown = GradingTemplatePeriod.objects.create(template=self.template, code="LEGACY1", name="Period 1", sequence_no=1)
        conflict = GradingTemplatePeriod.objects.create(template=self.template, code="MIDTERM", name="Final", sequence_no=2)
        self.assertIsNone(logical_period(unknown))
        self.assertIsNone(logical_period(conflict))
        other_template = GradingTemplate.objects.create(tenant=self.tenant, code="OLD", name="Historical")
        other_period = GradingTemplatePeriod.objects.create(template=other_template, code="PRELIM", name="Prelim", sequence_no=1, is_active=False)
        student_id = StudentPeriodGrade.objects.first().student_id
        for p in (unknown, conflict, other_period):
            StudentPeriodGrade.objects.create(tenant=self.tenant, campus=self.campus, offering=self.offering, template_period=p, student_id=student_id, period_grade=90)
        self.client.force_login(self.admin_user)
        response = self.search()
        self.assertEqual(len(response.context["rows"]), 2)
        self.assertContains(response, "2 saved physical period mapping(s) are unresolved")
        self.assertNotContains(response, "LEGACY1")

    def test_statuses_zero_and_missing_are_separate(self):
        grades = list(StudentPeriodGrade.objects.filter(offering=self.offering).order_by("id"))
        grades[0].period_grade = 0; grades[0].is_finalized = True; grades[0].save()
        grades[1].period_grade = None; grades[1].save()
        self.client.force_login(self.admin_user)
        row = self.search().context["rows"][0]
        self.assertEqual(row["graded_count"], 9)
        self.assertIsNone(row["minimum"])
        self.assertTrue(row["privacy_suppressed"])
        self.assertEqual(row["missing_count"], 1)
        self.assertEqual(row["status"], "Mixed")
        self.assertIsNone(row["average"])
        self.assertEqual(row["privacy_reason"], "complementary_population")
        final = self.search(status="finalized").context["rows"][0]
        self.assertEqual(final["graded_count"], 1)
        self.assertIsNone(final["average"])
        self.assertTrue(final["privacy_suppressed"])
        self.assertEqual(final["status_excluded_count"], 8)
        self.assertEqual(final["missing_count"], 1)
        self.assertEqual(self.search(status="unfinalized").context["rows"][0]["graded_count"], 8)

    def test_numeric_threshold_boundaries_and_perfect_concentration(self):
        values = list(map(Decimal, ("100", "90", "89.99", "75", "74.99", "70", "69.99", "60", "59.99", "0")))
        bands = distribution(values, Decimal("75"))
        self.assertEqual([b["count"] for b in bands], [2, 2, 2, 2, 2])
        self.assertEqual([b["range"] for b in bands], ["90 and above", "75 to below 90", "70 to below 75", "60 to below 70", "below 60"])
        StudentPeriodGrade.objects.filter(offering=self.offering).update(period_grade=100)
        self.client.force_login(self.admin_user)
        row = self.search().context["rows"][0]
        self.assertEqual(row["exact_100_count"], 10)
        self.assertEqual(row["exact_100_percentage"], 100)
        self.assertTrue(row["all_100"])

    def test_profile_template_tenant_and_system_threshold_precedence(self):
        self.client.force_login(self.admin_user)
        SystemSettingService.set("PASSING_GRADE_THRESHOLD", "80", tenant_id=self.tenant.id)
        self.assertEqual(self.search().context["rows"][0]["passing_threshold"], 75)
        profile = TenantGradingProfile.objects.create(tenant=self.tenant, campus=self.campus, profile_code="CUSTOM", profile_name="Custom", grading_template=self.template, passing_grade_threshold=82)
        row = self.search().context["rows"][0]
        self.assertEqual(row["passing_threshold"], 82)
        self.assertIn("profile", row["threshold_source"])
        profile.is_active = False; profile.save()
        self.template.is_published = False; self.template.save()
        self.assertEqual(self.search().context["rows"][0]["passing_threshold"], 80)
        SystemSettingService.set("PASSING_GRADE_THRESHOLD", "invalid", tenant_id=self.tenant.id)
        self.assertEqual(self.search().context["rows"][0]["passing_threshold"], 75)

    def test_faculty_labels_sorted_duplicate_and_missing_names(self):
        self.faculty_user.first_name = "Ana"; self.faculty_user.last_name = "Zulu"; self.faculty_user.save()
        other = self._user("duplicate", "Ana", "Zulu", self.tenant, self.campus, self.department)
        self._assign_faculty_role(other)
        FacultyAssignment.objects.create(offering=self.offering, faculty_user=other, response_status="ACCEPTED")
        blank = self._user("blank@example.com", "", "", self.tenant, self.campus, self.department)
        self._assign_faculty_role(blank)
        FacultyAssignment.objects.create(offering=self.offering, faculty_user=blank, response_status="ACCEPTED")
        self.client.force_login(self.admin_user)
        faculty = self.client.get(self.url).context["filter_options"]["faculty"]
        self.assertEqual(faculty[0]["label"], f"Faculty #{blank.id}")
        self.assertEqual(faculty[1]["label"], f"Zulu, Ana (#{self.faculty_user.id})")
        self.assertNotIn("@", str(faculty))

    def test_privacy_html_json_xlsx_have_no_student_identities(self):
        self.client.force_login(self.admin_user)
        response = self.search()
        detail = self.detail(response)
        export = self.client.get(self.url + "?" + response.context["export_query"])
        self.assertEqual(detail.status_code, 200)
        self.assertEqual(export.status_code, 200)
        for content in (response.content.decode(), detail.content.decode(), self._export_text(export)):
            for forbidden in ("ACC101-001", "Student1", "student_id", "student_no", "masked_name", "masked_student", "@example.com"):
                self.assertNotIn(forbidden, content)
        payload = detail.json()
        self.assertEqual(sum(f["count"] for f in payload["frequency"]), payload["graded_count"])
        self.assertIn("Computed Grade Component", self._export_text(export))

    def _privacy_assert_payload(self, payload):
        self.assertTrue(payload["privacy_suppressed"])
        for key in ("average", "minimum", "maximum", "exact_100_count", "exact_100_percentage", "all_100"):
            self.assertIsNone(payload[key], key)
        self.assertEqual(payload["frequency"], [])
        self.assertEqual(len(payload["high_grade_breakdown"]), 8)
        self.assertTrue(all(item["count"] is None and item["percentage"] is None
                            for item in payload["high_grade_breakdown"]))
        self.assertNotIn("equivalent_to", payload)
        self.assertNotIn("equivalent_dataset_id", payload)
        self.assertNotIn("equivalent_context", payload)
        self.assertTrue(all(b["count"] is None and b["percentage"] is None for b in payload["distribution"]))
        self.assertNotIn("student_id", payload)

    def _workbook(self, response):
        from io import BytesIO
        from openpyxl import load_workbook
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response["Content-Type"], "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
        return load_workbook(BytesIO(response.content))

    def _export_text(self, response):
        return str({sheet.title: list(sheet.values) for sheet in self._workbook(response)})

    def _privacy_xlsx_row(self, response, component="Periodic Examination"):
        records = self._worksheet_records(self._workbook(response), "Grade Distribution")
        record = next(r for r in records if r["Computed Grade Component"] == component)
        return {key: "" if value is None else str(value) for key, value in record.items()}

    def test_privacy_zero_one_to_four_five_and_six_html_json_xlsx_parity(self):
        self.client.force_login(self.admin_user)
        grades = list(StudentPeriodGrade.objects.filter(offering=self.offering).order_by("id"))
        for count in range(7):
            with self.subTest(count=count):
                StudentPeriodGrade.objects.filter(offering=self.offering).update(period_grade=None, exam_grade=None)
                StudentPeriodGrade.objects.filter(pk__in=[g.pk for g in grades[:count]]).update(period_grade=Decimal("83.27"), exam_grade=Decimal("83.27"))
                response = self.search()
                row = response.context["rows"][0]
                modal = self.detail(response).json()
                csv_response = self.client.get(self.url + "?" + response.context["export_query"])
                csv_row = self._privacy_xlsx_row(csv_response)
                self.assertEqual((row["graded_count"], row["missing_count"]), (count, 10 - count))
                self.assertEqual(row["analytics_state"], modal["analytics_state"])
                if 0 < count < 5:
                    self._privacy_assert_payload(row)
                    self._privacy_assert_payload(modal)
                    self.assertContains(response, "fewer than five students")
                    self.assertNotContains(response, "83.27")
                    self.assertNotIn("83.27", self._export_text(csv_response))
                    self.assertNotIn("83.27", str(modal))
                    self.assertNotContains(response, 'class="gdm-distribution"')
                    for grade in range(100, 94, -1):
                        self.assertEqual(csv_row[f"Exactly {grade}"], "")
                elif count == 0:
                    self.assertEqual(row["analytics_state"], "unavailable")
                    self.assertFalse(row["privacy_suppressed"])
                    self.assertIsNone(row["average"])
                else:
                    self.assertEqual(row["analytics_state"], "disclosed")
                    self.assertEqual(row["average"], Decimal("83.27"))
                    self.assertEqual(Decimal(modal["average"]), Decimal("83.27"))
                    self.assertTrue(all(csv_row[f"Exactly {grade}"] == "0" for grade in range(100, 94, -1)))
        StudentPeriodGrade.objects.filter(offering=self.offering).update(period_grade=None, exam_grade=None)
        StudentPeriodGrade.objects.filter(pk__in=[g.pk for g in grades[:5]]).update(period_grade=0, exam_grade=0)
        self.assertEqual(self.search().context["rows"][0]["average"], 0)

    def test_privacy_complementary_status_suppression_is_identical_across_endpoints(self):
        self.client.force_login(self.admin_user)
        ids = list(StudentPeriodGrade.objects.filter(offering=self.offering).order_by("id").values_list("id", flat=True))
        StudentPeriodGrade.objects.filter(offering=self.offering).update(period_grade=Decimal("87.43"), is_finalized=True)
        StudentPeriodGrade.objects.filter(pk=ids[-1]).update(is_finalized=False)
        for status in ("all", "finalized", "unfinalized"):
            response = self.search(status=status)
            row = response.context["rows"][0]
            self._privacy_assert_payload(row)
            self._privacy_assert_payload(self.detail(response).json())
            self.assertEqual(row["missing_count"], 0)
            for output in (response.content.decode(), self._export_text(self.client.get(self.url + "?" + response.context["export_query"]))):
                self.assertNotIn("87.43", output)
        StudentPeriodGrade.objects.filter(pk__in=ids[5:]).update(is_finalized=False)
        for status, count in (("all", 10), ("finalized", 5), ("unfinalized", 5)):
            row = self.search(status=status).context["rows"][0]
            self.assertEqual((row["analytics_state"], row["graded_count"]), ("disclosed", count))

    def test_privacy_saved_measures_and_component_distinct_populations(self):
        self.client.force_login(self.admin_user)
        ids = list(StudentPeriodGrade.objects.filter(offering=self.offering).order_by("id").values_list("student_id", flat=True))
        StudentPeriodGrade.objects.filter(student_id__in=ids[:3], offering=self.offering).update(exam_grade=Decimal("86.19"))
        StudentPeriodGrade.objects.filter(student_id__in=ids[:5], offering=self.offering).update(class_standing_grade=Decimal("78.31"))
        component = GradingTemplateComponent.objects.create(template_period=self.period, code="PRIV-EX", name="Private Exam", weight_percentage=100, is_exam_component=True, score_input_mode="DIRECT_PERCENTAGE")
        for i in range(2):
            activity = GradeActivity.objects.create(tenant=self.tenant, campus=self.campus, offering=self.offering,
                template_period=self.period, template_component=component, title=f"Private source {i}", total_score=100, created_by_user=self.faculty_user)
            for sid in ids[:3]:
                StudentActivityScore.objects.create(activity=activity, student_id=sid, raw_score=Decimal("82.17"), computed_score=Decimal("82.17"))
        data = self.detail(self.search()).json()
        self._privacy_assert_payload(data)
        self._privacy_assert_payload(data["measures"][0])
        self._privacy_assert_payload(data["measures"][1])
        self.assertEqual(data["measures"][1]["graded_count"], 5)
        for measure in data["measures"][2:]:
            self.assertEqual(measure["graded_count"], 3)
            self._privacy_assert_payload(measure)
        self.assertNotIn("86.19", str(data))
        self.assertNotIn("82.17", str(data))

    def _formula_fixture(self, *, supplemental_count=1, exam=True, supplemental_weight=20, common_weight=40, other_weight=40):
        """Persist engine-derived totals in the disposable test DB, not invented formulas."""
        ids = list(StudentPeriodGrade.objects.filter(offering=self.offering).order_by("id").values_list("student_id", flat=True))
        components = []
        for code, weight, values in (("Common", common_weight, list(range(70, 80))),
                                     ("Supplemental", supplemental_weight, [43] * supplemental_count),
                                     ("Other", other_weight, [80] * 10)):
            component = GradingTemplateComponent.objects.create(template_period=self.period, code=code,
                name=code, weight_percentage=weight, is_exam_component=exam if code != "Other" else not exam,
                score_input_mode="DIRECT_PERCENTAGE")
            activity = GradeActivity.objects.create(tenant=self.tenant, campus=self.campus, offering=self.offering,
                template_period=self.period, template_component=component, title=code, total_score=100,
                created_by_user=self.faculty_user)
            for sid, value in zip(ids, values):
                StudentActivityScore.objects.create(activity=activity, student_id=sid, raw_score=value, computed_score=value)
            components.append(component)
        self._persist_formula_totals(ids)
        self.client.force_login(self.admin_user)
        return ids, components

    def _persist_formula_totals(self, ids):
        for sid in ids:
            detail = FacultyGradingService.build_period_grade_detail_for_student(
                offering=self.offering, template_period=self.period, student_id=sid)
            StudentPeriodGrade.objects.filter(offering=self.offering, student_id=sid).update(
                period_grade=detail["period_grade"], exam_grade=detail["exam_grade"],
                class_standing_grade=detail["class_standing"])

    def test_formula_privacy_independent_public_payload_subtraction_attack(self):
        self._formula_fixture()
        self.assertEqual(sum(StudentPeriodGrade.objects.filter(offering=self.offering).values_list("exam_grade", flat=True)), 788)
        response = self.search()
        data = self.detail(response).json()
        common = next(m for m in data["measures"] if m["label"] == "Common")
        total = data["measures"][0]
        # Attacker uses only the actual public JSON, including chart/table frequencies.
        attacks = []
        if total["average"] is not None and common["average"] is not None:
            attacks.append(Decimal(total["average"]) * total["graded_count"] - Decimal(common["average"]) * common["graded_count"])
        if total["frequency"] and common["frequency"]:
            public_sum = lambda m: sum(Decimal(f["grade"]) * f["count"] for f in m["frequency"])
            attacks.append(public_sum(total) - public_sum(common))
        self.assertNotIn(Decimal("43"), attacks, "745/43/788 recovered from released fields")
        self._privacy_assert_payload(total)
        # A useful common breakdown remains public, but no related total can
        # complete the subtraction. This tests the actual retained disclosures.
        self.assertEqual(common["analytics_state"], "disclosed")
        self.assertEqual(sum(Decimal(f["grade"]) * f["count"] for f in common["frequency"]), 745)
        self.assertEqual(Decimal(common["average"]) * common["graded_count"], 745)
        self._privacy_assert_payload(data)

    def test_formula_privacy_html_json_xlsx_chart_consistency_and_statuses(self):
        ids, _ = self._formula_fixture()
        StudentPeriodGrade.objects.filter(offering=self.offering, student_id__in=ids[:5]).update(is_finalized=True)
        for status in ("all", "finalized", "unfinalized"):
            response = self.search(status=status)
            row = response.context["rows"][0]
            modal = self.detail(response).json()
            export = self.client.get(self.url + "?" + response.context["export_query"])
            csv_row = self._privacy_xlsx_row(export)
            self._privacy_assert_payload(row)
            self._privacy_assert_payload(modal)
            for grade in range(100, 94, -1):
                self.assertEqual(csv_row[f"Exactly {grade}"], "")
            self.assertNotContains(response, 'class="gdm-distribution"')
            for measure in modal["measures"]:
                if measure["analytics_state"] != "disclosed":
                    self.assertEqual(measure["frequency"], [])
                    self.assertIsNone(measure["average"])
                    self.assertTrue(all(b["count"] is None for b in measure["distribution"]))
                else:
                    self.assertGreaterEqual(measure["graded_count"], 5)
                    self.assertTrue(measure["frequency"])
            self.assertNotIn('"grade": "43', json.dumps(modal))
            self.assertEqual(row["missing_count"], 0)
            self.assertEqual(row["privacy_reason"], "formula_dependency")
        # A different physical class has no formula connection to this one.
        other = self._create_extra_offering("SAFE", "Unaffected", "SAFE-A")
        self._seed_grades(other, [90] * 10)
        safe = next(r for r in self.search().context["rows"] if r["offering_id"] == other.id)
        self.assertEqual((safe["analytics_state"], safe["average"]), ("disclosed", Decimal("90")))

    def test_formula_privacy_different_weights_period_and_bucket_equations(self):
        self._formula_fixture(common_weight=100, supplemental_weight=200, other_weight=100)
        rows = list(StudentPeriodGrade.objects.filter(offering=self.offering).values("period_grade", "exam_grade", "class_standing_grade"))
        self.assertEqual(sum(r["period_grade"] - r["exam_grade"] - r["class_standing_grade"] for r in rows), Decimal("43"))
        data = self.detail(self.search()).json()
        exam, standing = data["measures"][:2]
        attacks = []
        for field in ("average", "frequency"):
            def public_sum(m):
                if field == "average":
                    return Decimal(m[field]) * m["graded_count"] if m[field] is not None else None
                return sum(Decimal(f["grade"]) * f["count"] for f in m[field]) if m[field] else None
            totals = [public_sum(m) for m in (data, exam, standing)]
            if all(v is not None for v in totals):
                attacks.append(totals[0] - totals[1] - totals[2])
        # With weights 100/200/100, these three totals recover supplemental 43
        # even if BOTH individual exam components are withheld.
        self.assertNotIn(Decimal("43"), attacks)
        for measure in (data, exam):
            self._privacy_assert_payload(measure)
        self.assertEqual(Decimal(standing["average"]), 80)

    def test_formula_privacy_class_standing_activity_average_and_subtotal(self):
        ids, components = self._formula_fixture(supplemental_count=5, exam=False)
        activity = GradeActivity.objects.create(tenant=self.tenant, campus=self.campus, offering=self.offering,
            template_period=self.period, template_component=components[0], title="Small quiz", total_score=100,
            created_by_user=self.faculty_user)
        StudentActivityScore.objects.create(activity=activity, student_id=ids[0], raw_score=43, computed_score=43)
        self._persist_formula_totals(ids)
        data = self.detail(self.search()).json()
        # One quiz has ONE student. Withhold its dependent standing/period
        # totals while retaining the independently safe common activity.
        common = next(m for m in data["measures"] if m["label"] == "Common")
        self.assertEqual(common["graded_count"], 10)
        self.assertEqual(common["analytics_state"], "disclosed")
        self._privacy_assert_payload(next(m for m in data["measures"] if m["label"] == "Small quiz"))
        self._privacy_assert_payload(data["measures"][1])
        self._privacy_assert_payload(data)

    def test_formula_privacy_actual_percentage_configured_grade_differ(self):
        _, components = self._formula_fixture()
        activity = GradeActivity.objects.get(template_component=components[1])
        activity.total_score = 50
        activity.save()
        StudentActivityScore.objects.filter(activity=activity).update(raw_score=Decimal("21.5"), computed_score=Decimal("71.5"))
        self._persist_formula_totals(list(StudentPeriodGrade.objects.filter(offering=self.offering).values_list("student_id", flat=True)))
        data = self.detail(self.search()).json()
        self.assertFalse(any("Actual examination percentage" in m["label"] for m in data["measures"]))
        self._privacy_assert_payload(next(m for m in data["measures"] if m["label"] == "Supplemental"))
        self._privacy_assert_payload(data["measures"][0])
        self._privacy_assert_payload(data)

    def test_formula_privacy_weighted_subcomponent_details_with_small_activity(self):
        ids, components = self._formula_fixture(supplemental_count=5, exam=False)
        sub = GradingTemplateSubcomponent.objects.create(template_component=components[0], code="DETAILS",
            name="Weighted details", weight_percentage=100)
        dense = GradingTemplateDetail.objects.create(template_subcomponent=sub, code="DENSE", name="Dense", weight_percentage=70)
        sparse = GradingTemplateDetail.objects.create(template_subcomponent=sub, code="SPARSE", name="Sparse", weight_percentage=30)
        common = GradeActivity.objects.get(template_component=components[0])
        common.template_subcomponent = sub
        common.template_detail = dense
        common.save()
        activity = GradeActivity.objects.create(tenant=self.tenant, campus=self.campus, offering=self.offering,
            template_period=self.period, template_component=components[0], template_subcomponent=sub,
            template_detail=sparse, title="Private weighted detail", total_score=100, created_by_user=self.faculty_user)
        StudentActivityScore.objects.create(activity=activity, student_id=ids[0], raw_score=43, computed_score=43)
        self._persist_formula_totals(ids)
        response = self.search()
        data = self.detail(response).json()
        self._privacy_assert_payload(data)
        self._privacy_assert_payload(data["measures"][1])
        self.assertEqual(next(m for m in data["measures"] if m["label"] == "Common")["graded_count"], 10)

    def test_formula_privacy_sparse_attendance_session_protects_dense_standing(self):
        from apps.attendance.models import AttendanceRecord, AttendanceSession
        ids, components = self._formula_fixture(supplemental_count=5, exam=False)
        sub = GradingTemplateSubcomponent.objects.create(template_component=components[0], code="ATT",
            name="Attendance", weight_percentage=100, is_attendance_component=True)
        for index, participants in enumerate((ids, ids[:1])):
            session = AttendanceSession.objects.create(tenant=self.tenant, campus=self.campus, offering=self.offering,
                template_period=self.period, session_date=date(2025, 7, index + 1), created_by_user=self.faculty_user)
            for sid in participants:
                AttendanceRecord.objects.create(session=session, student_id=sid, status_code="P")
        self._persist_formula_totals(ids)
        data = self.detail(self.search()).json()
        self._privacy_assert_payload(data)
        self.assertEqual(next(m for m in data["measures"] if m["label"] == "Common")["analytics_state"], "disclosed")
        self._privacy_assert_payload(data["measures"][1])

    def test_formula_privacy_five_boundary_retains_safe_detailed_analytics(self):
        ids, _ = self._formula_fixture(supplemental_count=5)
        for status in ("all", "unfinalized"):
            data = self.detail(self.search(status=status)).json()
            self.assertEqual(data["analytics_state"], "disclosed")
            self.assertTrue(all(m["analytics_state"] == "disclosed" for m in data["measures"]))
            supplemental = next(m for m in data["measures"] if m["label"] == "Supplemental")
            self.assertEqual((supplemental["graded_count"], Decimal(supplemental["average"])), (5, Decimal("43")))
        score = StudentActivityScore.objects.get(student_id=ids[4], activity__title="Supplemental")
        score.delete()
        self._persist_formula_totals(ids)
        data = self.detail(self.search()).json()
        self._privacy_assert_payload(data)
        self._privacy_assert_payload(data["measures"][0])

    def test_formula_privacy_zero_weight_independence_preserves_period_and_standing(self):
        self._formula_fixture(common_weight=0, supplemental_weight=0, other_weight=100)
        response = self.search()
        row = response.context["rows"][0]
        data = self.detail(response).json()
        self.assertEqual((row["analytics_state"], row["average"]), ("disclosed", Decimal("80")))
        self.assertEqual(Decimal(data["average"]), row["average"])
        self._privacy_assert_payload(data["measures"][0])
        self.assertEqual(Decimal(data["measures"][1]["average"]), Decimal("80"))
        self.assertEqual(self._privacy_xlsx_row(self.client.get(self.url + "?" + response.context["export_query"]))["Exactly 100"], "")

    def test_privacy_overlap_guard_does_not_change_summary_when_opening_modal(self):
        ids = list(StudentPeriodGrade.objects.filter(offering=self.offering).order_by("id").values_list("student_id", flat=True))
        StudentPeriodGrade.objects.filter(student_id__in=ids[:9], offering=self.offering).update(exam_grade=Decimal("84.29"))
        self.client.force_login(self.admin_user)
        for status in ("all", "unfinalized"):
            response = self.search(status=status)
            row = response.context["rows"][0]
            data = self.detail(response).json()
            self.assertEqual(row["average"], data["average"])
            self._privacy_assert_payload(row)
            self._privacy_assert_payload(data)
            self._privacy_assert_payload(data["measures"][0])
            self.assertEqual(data["measures"][0]["privacy_reason"], "complementary_population")
            self.assertNotIn("84.29", str(data))
        # Changed status without matching signed filters cannot obtain another population.
        params = parse_qs(response.context["page_query"])
        params["status"] = "finalized"
        self.assertEqual(self.client.get(reverse("admin_portal:grade_distribution_monitor_detail", args=[self.offering.id, self.period.id]), params).status_code, 400)

    def test_privacy_historical_transfer_and_multi_campus_authorization(self):
        campus, dept, target = self._second_campus(grant_admin=True)
        self.offering.status = "ARCHIVED"
        self.offering.save()
        ids = list(StudentPeriodGrade.objects.filter(offering=self.offering).order_by("id").values_list("student_id", flat=True))
        StudentPeriodGrade.objects.filter(offering=self.offering).update(period_grade=None)
        StudentPeriodGrade.objects.filter(offering=self.offering, student_id__in=ids[:4]).update(period_grade=Decimal("88.13"))
        Student.objects.filter(pk=ids[0]).update(campus=campus, department=dept, program=target.program)
        self.client.force_login(self.admin_user)
        response = self.search(campus_id="all")
        self.assertEqual(len(response.context["rows"]), 2)
        for row in response.context["rows"]:
            self._privacy_assert_payload(row)
            self._privacy_assert_payload(self.detail(response, row=row).json())
        retained = next(r for r in response.context["rows"] if r["offering_id"] == self.offering.id)
        self.assertEqual(retained["eligible_count"], 10)
        self.assertNotIn("88.13", self._export_text(self.client.get(self.url + "?" + response.context["export_query"])))
        UserPermission.objects.create(user=self.admin_user, permission=Permission.objects.get(code="grade_distribution_monitor.read"), tenant=self.tenant, campus=campus, grant_type="DENY")
        self.assertEqual(self.client.get(self.url + "?" + response.context["export_query"]).status_code, 400)
        self.assertEqual(len(self.search(campus_id="all").context["rows"]), 1)

    def test_detail_and_exports_require_valid_context(self):
        self.client.force_login(self.admin_user)
        response = self.search()
        self.assertEqual(self.client.get(self.url, {"export": "csv"}).status_code, 400)
        row = response.context["rows"][0]
        detail_url = reverse("admin_portal:grade_distribution_monitor_detail", args=[row["offering_id"], row["physical_period_id"]])
        self.assertEqual(self.client.get(detail_url, {"apply": 1, "academic_year_id": self.academic_year.id, "term_id": self.term.id, "period": "PRELIM"}).status_code, 400)
        query = parse_qs(row["detail_query"])
        query["period"] = "FINAL"
        self.assertEqual(self.client.get(detail_url, {k: v[0] if isinstance(v, list) else v for k, v in query.items()}).status_code, 400)
        query = parse_qs(row["detail_query"]); query["context"] = "tampered"
        self.assertEqual(self.client.get(detail_url, {k: v[0] if isinstance(v, list) else v for k, v in query.items()}).status_code, 400)

    def test_expired_context(self):
        self.client.force_login(self.admin_user)
        response = self.search()
        with patch("django.core.signing.time.time", return_value=timezone.now().timestamp() + 4000):
            self.assertEqual(self.detail(response).status_code, 400)

    def test_direct_deny_and_revoked_supervision_block_detail_export(self):
        self.client.force_login(self.admin_user)
        response = self.search()
        deny = UserPermission.objects.create(user=self.admin_user, permission=Permission.objects.get(code="grade_distribution_monitor.read"), tenant=self.tenant, campus=self.campus, grant_type="DENY")
        self.assertEqual(self.detail(response).status_code, 403)
        self.assertEqual(self.client.get(self.url + "?" + response.context["export_query"]).status_code, 403)
        deny.delete()
        UserRole.objects.filter(user=self.faculty_user).update(is_active=False)
        self.assertEqual(self.detail(response).status_code, 400)
        self.assertEqual(self.client.get(self.url + "?" + response.context["export_query"]).status_code, 400)

    def test_historical_saved_values_retained_without_current_active_roster(self):
        self.offering.status = "ARCHIVED"; self.offering.is_active = False; self.offering.save()
        FacultyAssignment.objects.filter(offering=self.offering).update(is_active=False)
        Enrollment.objects.filter(course_offering=self.offering).update(is_active=False)
        Student.objects.filter(tenant=self.tenant).update(is_active=False)
        self.academic_year.is_active = False; self.academic_year.save()
        self.term.is_active = False; self.term.save()
        self.client.force_login(self.admin_user)
        response = self.search()
        row = response.context["rows"][0]
        self.assertEqual(row["graded_count"], 10)
        self.assertEqual(row["average"], Decimal("90.40"))
        self.assertTrue(row["historical"])
        payload = self.detail(response).json()
        self.assertEqual(len(payload["measures"]), 2)
        self.assertIsNone(payload["incomplete_source_count"])
        self.assertIn("Original roster", " ".join(payload["notes"]))

    def test_historical_transfer_preserves_summary_modal_xlsx_and_original_authority(self):
        self.offering.status = "ARCHIVED"
        self.offering.save()
        campus, dept, target = self._second_campus()
        student = StudentPeriodGrade.objects.filter(offering=self.offering, period_grade=95).get().student
        old_enrollment = list(Enrollment.objects.filter(student=student).values())
        self.client.force_login(self.admin_user)
        keys = ("eligible_count", "graded_count", "average", "distribution", "missing_count", "coverage_rate",
                "finalized_count", "unfinalized_count", "status_excluded_count", "physical_period_id")
        for finalized in (False, True):
            with self.subTest(finalized=finalized):
                Student.objects.filter(pk=student.pk).update(campus=self.campus, department=self.department, program=self.program)
                StudentPeriodGrade.objects.filter(offering=self.offering).update(is_finalized=finalized)
                saved_before = list(StudentPeriodGrade.objects.filter(offering=self.offering).values())
                captures = []
                for status in ("all", "finalized", "unfinalized"):
                    before = self.search(status=status)
                    row = before.context["rows"][0]
                    modal = self.detail(before).json()
                    modal.pop("detail_query")  # Signing timestamp is not an analytics value.
                    captures.append((status, {k: row[k] for k in keys}, modal,
                                     self._export_text(self.client.get(self.url + "?" + before.context["export_query"]))))
                Student.objects.filter(pk=student.pk).update(campus=campus, department=dept, program=target.program)
                for status, summary, modal, csv in captures:
                    after = self.search(status=status)
                    self.assertEqual({k: after.context["rows"][0][k] for k in keys}, summary)
                    after_modal = self.detail(after).json()
                    after_modal.pop("detail_query")
                    self.assertEqual(after_modal, modal)
                    self.assertEqual(self._export_text(self.client.get(self.url + "?" + after.context["export_query"])), csv)
                row = self.search().context["rows"][0]
                self.assertEqual((row["eligible_count"], row["graded_count"], row["average"], row["coverage_rate"]),
                                 (10, 10, Decimal("90.40"), Decimal("100.00")))
                self.assertEqual([b["count"] for b in row["distribution"]], [8, 2, 0, 0, 0])
                self.assertEqual(list(Enrollment.objects.filter(student=student).values()), old_enrollment)
                self.assertEqual(list(StudentPeriodGrade.objects.filter(offering=self.offering).values()), saved_before)
        authorized = self.search()
        head = self._user("new_head", "New", "Head", self.tenant, campus, dept)
        self._assign_permissions(head, include_monitor=True)
        self.client.force_login(head)
        own = self.search()
        self.assertEqual([r["offering_id"] for r in own.context["rows"]], [target.id])
        previous_url = reverse("admin_portal:grade_distribution_monitor_detail", args=[self.offering.id, self.period.id])
        self.assertEqual(self.client.get(previous_url + "?" + own.context["page_query"]).status_code, 404)
        self.assertNotIn("ACC101", self._export_text(self.client.get(self.url + "?" + own.context["export_query"])))
        self.assertEqual(self.detail(authorized).status_code, 400)
        self.assertEqual(self.client.get(self.url + "?" + authorized.context["export_query"]).status_code, 400)
        self.assertTrue(self.search(campus_id=self.campus.id).context["errors"])
        self.client.force_login(self.admin_user)
        UserPermission.objects.create(user=self.admin_user, permission=Permission.objects.get(code="grade_distribution_monitor.read"),
                                      tenant=self.tenant, campus=self.campus, grant_type="DENY")
        self.assertEqual(self.detail(authorized).status_code, 403)
        self.assertEqual(self.client.get(self.url + "?" + authorized.context["export_query"]).status_code, 403)
        UserPermission.objects.filter(user=self.admin_user).delete()
        UserRole.objects.filter(user=self.faculty_user, campus=self.campus).update(is_active=False)
        self.assertEqual(self.detail(authorized).status_code, 400)
        self.assertEqual(self.client.get(self.url + "?" + authorized.context["export_query"]).status_code, 400)
        self.assertEqual(self.search().context["rows"], [])

    def test_historical_transfer_does_not_override_tenant_or_retained_enrollment_evidence(self):
        self.offering.status = "ARCHIVED"
        self.offering.save()
        student = StudentPeriodGrade.objects.filter(offering=self.offering, period_grade=95).get().student
        foreign = Tenant.objects.create(code="FOREIGN-HISTORY", name="Foreign history")
        Student.objects.filter(pk=student.pk).update(tenant=foreign)
        self.client.force_login(self.admin_user)
        row = self.search().context["rows"][0]
        self.assertEqual((row["eligible_count"], row["graded_count"], row["average"]), (9, 9, Decimal("89.89")))
        Student.objects.filter(pk=student.pk).update(tenant=self.tenant)
        Enrollment.objects.filter(student=student).delete()
        row = self.search().context["rows"][0]
        self.assertEqual((row["eligible_count"], row["graded_count"]), (9, 9))

    def test_scope_wide_class_limit_rejects_oversized_later_page_without_partial_results(self):
        for i in range(26):
            later = self._create_extra_offering(f"GLOBAL{i}", "Global limit", f"GL{i}")
            self._seed_grades(later, [90])
        self.client.force_login(self.admin_user)
        with patch.object(GradeDistributionMonitorService, "MAX_CLASS_STUDENTS", 10):
            valid = self.search()
            self.assertEqual((len(valid.context["rows"]), valid.context["result_count"]), (25, 27))
            for index in range(2, 12):
                student = Student.objects.create(tenant=self.tenant, campus=self.campus, department=self.department,
                    program=self.program, student_no=f"GLOBAL-LAST-{index}", first_name="Limit", last_name="Student")
                Enrollment.objects.create(tenant=self.tenant, campus=self.campus, academic_year=self.academic_year,
                    term=self.term, student=student, course_offering=later)
            requests = ({"apply": "1", "academic_year_id": self.academic_year.id, "term_id": self.term.id, "period": "PRELIM"},
                        parse_qs(valid.context["page_query"] + "&page=2"), parse_qs(valid.context["export_query"]))
            for params in requests:
                with CaptureQueriesContext(connection) as queries:
                    response = self.client.get(self.url, params)
                if "export" in params:
                    self.assertEqual(response.status_code, 400)
                    self.assertFalse(response.has_header("Content-Disposition"))
                else:
                    self.assertEqual(response.context["rows"], [])
                    self.assertIsNone(response.context["page_obj"])
                    self.assertFalse(response.context["searched"])
                self.assertContains(response, "10 student analysis limit", status_code=response.status_code)
                self.assertContains(response, "no records were truncated", status_code=response.status_code)
                aggregate = [q["sql"] for q in queries if 'enrollments' in q["sql"] and 'GROUP BY' in q["sql"] and 'COUNT(DISTINCT' in q["sql"]]
                self.assertEqual(len(aggregate), 1)
                self.assertFalse(any('student_period_grades' in q["sql"] and 'LIMIT 25' in q["sql"] for q in queries))
            detail = self.detail(valid)
            self.assertEqual(detail.status_code, 400)
            self.assertIn("10 student analysis limit", " ".join(detail.json()["errors"]))

    def test_global_limit_uses_historical_roster_and_live_eligibility(self):
        self.client.force_login(self.admin_user)
        Student.objects.filter(pk=StudentPeriodGrade.objects.filter(offering=self.offering).first().student_id).update(is_active=False)
        with patch.object(GradeDistributionMonitorService, "MAX_CLASS_STUDENTS", 9):
            live = self.search()
            self.assertFalse(live.context["errors"])
            self.assertEqual(live.context["rows"][0]["eligible_count"], 9)
            self.period.is_active = False
            self.period.save()
            self.assertTrue(self.search().context["errors"])
            self.period.is_active = True
            self.period.save()
            self.offering.status = "ARCHIVED"
            self.offering.save()
            self.assertTrue(self.search().context["errors"])
            with patch.object(GradeDistributionMonitorService, "MAX_CLASS_STUDENTS", 10):
                self.assertEqual(self.search().context["rows"][0]["eligible_count"], 10)

    def test_real_two_thousand_student_limit_accepts_boundary_and_rejects_next(self):
        students = Student.objects.bulk_create([
            Student(tenant=self.tenant, campus=self.campus, department=self.department, program=self.program,
                    student_no=f"REAL-LIMIT-{i}", first_name="Boundary", last_name="Student") for i in range(1991)
        ])
        enrollments = [Enrollment(tenant=self.tenant, campus=self.campus, academic_year=self.academic_year, term=self.term,
                                 course_offering=self.offering, student=s) for s in students]
        Enrollment.objects.bulk_create(enrollments[:1990])
        self.client.force_login(self.admin_user)
        exact = self.search()
        self.assertFalse(exact.context["errors"])
        self.assertEqual(exact.context["rows"][0]["eligible_count"], 2000)
        self.assertEqual(self.detail(exact).status_code, 200)
        self.assertEqual(self.client.get(self.url + "?" + exact.context["export_query"]).status_code, 200)
        Enrollment.objects.bulk_create(enrollments[1990:])
        rejected = self.search()
        self.assertEqual(rejected.context["rows"], [])
        self.assertContains(rejected, "2,000 student analysis limit")
        self.assertEqual(self.detail(exact).status_code, 400)
        self.assertEqual(self.client.get(self.url + "?" + exact.context["export_query"]).status_code, 400)

    def test_real_scope_limits_500_rows_and_50000_saved_and_roster_records(self):
        # Reuse one hundred students across five hundred distinct classes, all in scope.
        students = list(Student.objects.filter(tenant=self.tenant).order_by("id"))
        students += Student.objects.bulk_create([
            Student(tenant=self.tenant, campus=self.campus, department=self.department, program=self.program,
                    student_no=f"VOLUME-{i}", first_name="Volume", last_name="Student") for i in range(90)
        ])
        sections = Section.objects.bulk_create([
            Section(tenant=self.tenant, campus=self.campus, department=self.department, program=self.program,
                    code=f"VOLUME-{i}", name=f"Volume {i}") for i in range(499)
        ])
        offerings = [self.offering] + CourseOffering.objects.bulk_create([
            CourseOffering(tenant=self.tenant, campus=self.campus, department=self.department, program=self.program,
                academic_year=self.academic_year, term=self.term, course=self.course, section=section) for section in sections
        ])
        FacultyAssignment.objects.bulk_create([
            FacultyAssignment(tenant=self.tenant, campus=self.campus, offering=o, faculty_user=self.faculty_user,
                              response_status="ACCEPTED", accepted_at=timezone.now()) for o in offerings[1:]
        ])
        pairs = [(o, student) for o in offerings for student in students
                 if o.pk != self.offering.pk or student.pk not in {s.pk for s in students[:10]}]
        Enrollment.objects.bulk_create([
            Enrollment(tenant=self.tenant, campus=self.campus, academic_year=self.academic_year, term=self.term,
                       course_offering=o, student=student) for o, student in pairs
        ], batch_size=500)
        StudentPeriodGrade.objects.bulk_create([
            StudentPeriodGrade(tenant=self.tenant, campus=self.campus, offering=o, template_period=self.period,
                               student=student, period_grade=90) for o, student in pairs
        ], batch_size=500)
        self.assertEqual(Enrollment.objects.count(), 50000)
        self.assertEqual(StudentPeriodGrade.objects.count(), 50000)
        self.client.force_login(self.admin_user)
        import time
        started = time.perf_counter()
        exact = self.search()
        page_one_seconds = time.perf_counter() - started
        self.assertFalse(exact.context["errors"])
        self.assertEqual((exact.context["result_count"], len(exact.context["rows"])), (500, 25))
        self.assertEqual(self.detail(exact).status_code, 200)
        started = time.perf_counter()
        self.assertEqual(self.client.get(self.url + "?" + exact.context["page_query"] + "&page=20").status_code, 200)
        page_twenty_seconds = time.perf_counter() - started
        started = time.perf_counter()
        csv = self.client.get(self.url + "?" + exact.context["export_query"])
        export_seconds = time.perf_counter() - started
        print(f"GRADE100_BOUNDARY: 500 classes / 50,000 saved+roster rows: page1={page_one_seconds:.3f}s page20={page_twenty_seconds:.3f}s export={export_seconds:.3f}s")
        self.assertEqual(csv.status_code, 200)
        workbook = self._workbook(csv)
        self.assertEqual(workbook["Grade Distribution"].max_row, 1501)

        def assert_rejected(message):
            for params in ({"apply": "1", "academic_year_id": self.academic_year.id, "term_id": self.term.id, "period": "PRELIM"},
                           parse_qs(exact.context["page_query"] + "&page=20")):
                rejected = self.client.get(self.url, params)
                self.assertEqual(rejected.context["rows"], [])
                self.assertContains(rejected, message)
            detail = self.detail(exact)
            self.assertEqual(detail.status_code, 400)
            self.assertIn(message, " ".join(detail.json()["errors"]))
            export = self.client.get(self.url + "?" + exact.context["export_query"])
            self.assertEqual(export.status_code, 400)
            self.assertFalse(export.has_header("Content-Disposition"))

        extra_period = GradingTemplatePeriod.objects.create(template=self.template, code="PrelimExam", name="Prelim", sequence_no=5)
        extra_grade = StudentPeriodGrade.objects.create(tenant=self.tenant, campus=self.campus, offering=self.offering,
            template_period=extra_period, student=students[0], period_grade=90)
        assert_rejected("50,000 saved grade records")
        extra_grade.delete()
        new_student = Student.objects.create(tenant=self.tenant, campus=self.campus, department=self.department,
            program=self.program, student_no="VOLUME-EXCESS", first_name="Excess", last_name="Student")
        extra_enrollment = Enrollment.objects.create(tenant=self.tenant, campus=self.campus, academic_year=self.academic_year,
            term=self.term, course_offering=self.offering, student=new_student)
        assert_rejected("50,000 records")
        extra_enrollment.delete()
        StudentPeriodGrade.objects.filter(offering=self.offering, student=students[-1], template_period=self.period).delete()
        StudentPeriodGrade.objects.create(tenant=self.tenant, campus=self.campus, offering=self.offering,
            template_period=extra_period, student=students[0], period_grade=90)
        self.assertEqual(StudentPeriodGrade.objects.count(), 50000)
        assert_rejected("500 class/period rows")

    def test_all_search_limits_accept_exact_boundary_and_reject_one_over(self):
        self.client.force_login(self.admin_user)
        self.assertEqual((GradeDistributionMonitorService.MAX_CLASS_ROWS, GradeDistributionMonitorService.MAX_SAVED_ROWS,
                          GradeDistributionMonitorService.MAX_ROSTER_ROWS, GradeDistributionMonitorService.MAX_CLASS_STUDENTS),
                         (500, 50000, 50000, 2000))
        # Scale only constants: the same database predicates and comparisons execute.
        for name, exact in (("MAX_CLASS_ROWS", 1), ("MAX_SAVED_ROWS", 10), ("MAX_ROSTER_ROWS", 10), ("MAX_CLASS_STUDENTS", 10)):
            with self.subTest(limit=name), patch.object(GradeDistributionMonitorService, name, exact):
                response = self.search()
                self.assertFalse(response.context["errors"])
                self.assertEqual(len(response.context["rows"]), 1)
                self.assertEqual(self.detail(response).status_code, 200)
                self.assertEqual(self.client.get(self.url + "?" + response.context["export_query"]).status_code, 200)
                with patch.object(GradeDistributionMonitorService, name, exact - 1):
                    for params in ({"apply": "1", "academic_year_id": self.academic_year.id, "term_id": self.term.id, "period": "PRELIM"},
                                   parse_qs(response.context["page_query"] + "&page=1")):
                        rejected = self.client.get(self.url, params)
                        self.assertTrue(rejected.context["errors"])
                        self.assertEqual(rejected.context["rows"], [])
                    self.assertEqual(self.detail(response).status_code, 400)
                    self.assertEqual(self.client.get(self.url + "?" + response.context["export_query"]).status_code, 400)

    def test_summary_uses_simple_gradebook_labels_and_excel_link(self):
        self.client.force_login(self.admin_user)
        response = self.search()
        for text in ("Period Grades", "Examination Grades", "Class Standing Grades", "Average Grade", "Export Excel", "Base 50"):
            self.assertContains(response, text)
        self.assertIn("export=xlsx", response.context["export_query"])
        for text in ("Threshold-Aware Distribution", "Class Ave.", "Export CSV"):
            self.assertNotContains(response, text)

    def test_historical_assignment_without_acceptance_evidence_unavailable(self):
        self.offering.status = "CLOSED"; self.offering.save()
        FacultyAssignment.objects.filter(offering=self.offering).update(is_active=False, accepted_at=None)
        self.client.force_login(self.admin_user)
        self.assertEqual(self.search().context["rows"], [])

    def test_live_roster_excludes_inactive_dropped_and_foreign_students(self):
        first = Enrollment.objects.filter(course_offering=self.offering).first()
        first.is_active = False; first.save()
        second = Enrollment.objects.filter(course_offering=self.offering, is_active=True).first()
        second.enrollment_status = "DRP"; second.save()
        self.client.force_login(self.admin_user)
        row = self.search().context["rows"][0]
        self.assertEqual(row["eligible_count"], 8)
        self.assertEqual(row["graded_count"], 8)

    def test_database_pagination_and_export_limits(self):
        for i in range(26):
            offering = self._create_extra_offering(f"PAGE{i}", "Pagination Course", f"SEC{i}")
            self._seed_grades(offering, [90])
        self.client.force_login(self.admin_user)
        with CaptureQueriesContext(connection) as queries:
            response = self.search()
        self.assertEqual(len(response.context["rows"]), 25)
        self.assertEqual(response.context["result_count"], 27)
        self.assertFalse(any('GROUP BY' in q["sql"] and 'LIMIT 25' in q["sql"] for q in queries))
        page_two = self.client.get(self.url + "?" + response.context["page_query"] + "&page=2")
        self.assertEqual(len(page_two.context["rows"]), 2)
        self.assertEqual(self.client.get(self.url + "?" + response.context["export_query"]).status_code, 200)
        with patch.object(GradeDistributionMonitorService, "MAX_CLASS_ROWS", 26):
            narrowed = self.search()
            self.assertEqual(narrowed.context["rows"], [])
            self.assertContains(narrowed, "Narrow your criteria")
            self.assertEqual(self.client.get(self.url + "?" + response.context["export_query"]).status_code, 400)
        with patch.object(GradeDistributionMonitorService, "MAX_SAVED_ROWS", 5):
            self.assertTrue(self.search().context["errors"])
        with patch.object(GradeDistributionMonitorService, "MAX_CLASS_STUDENTS", 5):
            self.assertTrue(self.search().context["errors"])

    def test_monitor_never_writes_gradebook_or_calls_recompute(self):
        self.client.force_login(self.admin_user)
        before = list(StudentPeriodGrade.objects.values())
        with patch.object(FacultyGradingService, "recompute_period_summary", side_effect=AssertionError("recompute forbidden")):
            with CaptureQueriesContext(connection) as queries:
                response = self.search()
                self.detail(response)
                self.client.get(self.url + "?" + response.context["export_query"])
        self.assertEqual(list(StudentPeriodGrade.objects.values()), before)
        writes = [q["sql"] for q in queries if q["sql"].lstrip().upper().startswith(("INSERT", "UPDATE", "DELETE"))]
        self.assertFalse(any('student_period_grades' in q or 'student_activity_scores' in q or 'grade_submissions' in q or 'audit_logs' in q for q in writes))

    def test_raw_examination_and_configured_activity_grade_remain_distinct(self):
        exam = GradingTemplateComponent.objects.create(template_period=self.period, code="EXAM", name="Exam", weight_percentage=50, is_exam_component=True)
        standing = GradingTemplateComponent.objects.create(template_period=self.period, code="CS", name="Class Standing", weight_percentage=50)
        activity = GradeActivity.objects.create(tenant=self.tenant, campus=self.campus, offering=self.offering, template_period=self.period, template_component=exam, title="Exam", total_score=50, created_by_user=self.faculty_user)
        quiz = GradeActivity.objects.create(tenant=self.tenant, campus=self.campus, offering=self.offering, template_period=self.period, template_component=standing, title="Quiz", total_score=100, created_by_user=self.faculty_user)
        students = list(Student.objects.filter(enrollments__course_offering=self.offering).order_by("id")[:5])
        for student in students:
            StudentActivityScore.objects.create(activity=activity, student=student, raw_score=25, computed_score=75)
            StudentActivityScore.objects.create(activity=quiz, student=student, raw_score=100, computed_score=100)
        StudentPeriodGrade.objects.filter(student__in=students, offering=self.offering).update(exam_grade=75, class_standing_grade=100)
        self.client.force_login(self.admin_user)
        data = self.detail(self.search()).json()
        measures = {m["label"]: m for m in data["measures"]}
        self.assertNotIn("Exam - Actual examination percentage", measures)
        self.assertEqual(set(StudentActivityScore.objects.filter(activity=activity).values_list("raw_score", flat=True)), {Decimal("25")})
        self.assertEqual(Decimal(measures["Exam"]["average"]), 75)
        self.assertEqual(Decimal(measures["Quiz"]["average"]), 100)
        self.assertNotIn("Examination - Exam", measures)
        self.assertEqual(data["incomplete_source_count"], 5)

    def test_direct_percentage_multiple_components_and_no_combined_normalization(self):
        students = list(Student.objects.filter(enrollments__course_offering=self.offering).order_by("id")[:5])
        for i in range(2):
            c = GradingTemplateComponent.objects.create(template_period=self.period, code=f"EX{i}", name=f"Exam {i}", weight_percentage=50, is_exam_component=True, score_input_mode="DIRECT_PERCENTAGE")
            a = GradeActivity.objects.create(tenant=self.tenant, campus=self.campus, offering=self.offering, template_period=self.period, template_component=c, title=f"Exam {i}", total_score=100, created_by_user=self.faculty_user)
            for student in students:
                StudentActivityScore.objects.create(activity=a, student=student, raw_score=100, computed_score=100)
        StudentPeriodGrade.objects.filter(student__in=students, offering=self.offering).update(exam_grade=200)
        self.client.force_login(self.admin_user)
        data = self.detail(self.search()).json()
        total = data["measures"][0]
        self.assertEqual(Decimal(total["average"]), 200)
        self.assertEqual(total["exact_100_count"], 0)
        components = [m for m in data["measures"] if m["label"].startswith("Exam ")]
        self.assertEqual(len(components), 2)
        self.assertTrue(all(m["all_100"] for m in components))

    def test_monitor_scope_follows_faculty_assignment_scope_not_offering_department(self):
        reviewer_department = Department.objects.create(
            tenant=self.tenant,
            campus=self.campus,
            code="FVW_COLL_IS",
            name="Fairview Information Systems",
        )
        scoped_reviewer = self._user(
            "ac_scope",
            "Area",
            "Chair",
            self.tenant,
            self.campus,
            reviewer_department,
        )
        scoped_faculty = self._user(
            "faculty_scope",
            "Scoped",
            "Faculty",
            self.tenant,
            self.campus,
            reviewer_department,
        )
        self._assign_permissions(scoped_reviewer, include_monitor=True)
        self._assign_faculty_role(scoped_faculty)
        offering = self._create_extra_offering("MISMATCH101", "Offering Department Mismatch", "BSA1D")
        FacultyAssignment.objects.create(
            tenant=self.tenant,
            campus=self.campus,
            offering=offering,
            faculty_user=scoped_faculty,
            response_status=FacultyAssignment.ResponseStatus.ACCEPTED,
            accepted_at=timezone.now(),
            is_primary=True,
        )
        self._seed_grades(offering, [91, 92, 93, 94, 95, 96, 97, 98, 81, 76])

        self.client.force_login(scoped_reviewer)
        response = self.search(department_id=reviewer_department.id, faculty_id=scoped_faculty.id)

        self.assertEqual(response.status_code, 200)
        self.assertContains(response, "MISMATCH101")
        self.assertContains(response, "Offering Department Mismatch")

        self.assertIn(scoped_faculty.id, [f["id"] for f in response.context["filter_options"]["faculty"]])
        self.assertIn(reviewer_department.id, [d["id"] for d in response.context["filter_options"]["departments"]])
        self.assertEqual(len(response.context["rows"]), 1)
        row = response.context["rows"][0]
        detail = self.detail(response, row)
        self.assertEqual(detail.status_code, 200)
        self.assertEqual(detail.json()["course_code"], "MISMATCH101")
        export = self.client.get(self.url + "?" + response.context["export_query"])
        self.assertIn("MISMATCH101", self._export_text(export))
        self.assertNotIn("ACC101", self._export_text(export))

    def _second_campus(self, grant_admin=False):
        campus = Campus.objects.create(tenant=self.tenant, code="CUB", name="Cubao")
        dept = Department.objects.create(tenant=self.tenant, campus=campus, code="BSA-CUB", name="Accountancy")
        program = Program.objects.create(tenant=self.tenant, campus=campus, department=dept, code="BSA", name="BSA")
        course = Course.objects.create(tenant=self.tenant, campus=campus, department=dept, code="CUB101", title="Cubao Accounting")
        section = Section.objects.create(tenant=self.tenant, campus=campus, department=dept, program=program, code="CUB1", name="CUB1")
        offering = CourseOffering.objects.create(tenant=self.tenant, campus=campus, department=dept, program=program, academic_year=self.academic_year, term=self.term, course=course, section=section)
        UserRole.objects.create(user=self.faculty_user, role=Role.objects.get(code="FACULTY"), tenant=self.tenant, campus=campus, department=dept)
        FacultyAssignment.objects.create(tenant=self.tenant, campus=campus, offering=offering, faculty_user=self.faculty_user, response_status="ACCEPTED", accepted_at=timezone.now())
        self._seed_grades(offering, [100, 100])
        if grant_admin:
            role = self.admin_user.user_roles.exclude(role__code="FACULTY").get().role
            UserRole.objects.create(user=self.admin_user, role=role, tenant=self.tenant, campus=campus, department=dept)
        return campus, dept, offering

    def test_same_faculty_in_other_campus_cannot_leak_filters_results_modal_xlsx(self):
        campus, dept, offering = self._second_campus()
        self.client.force_login(self.admin_user)
        response = self.search(campus_id="all")
        self.assertEqual(len(response.context["rows"]), 1)
        self.assertNotContains(response, "Cubao Accounting")
        self.assertNotIn(dept.id, [d["id"] for d in response.context["filter_options"]["departments"]])
        self.assertNotIn("CUB101", self._export_text(self.client.get(self.url + "?" + response.context["export_query"])))
        url = reverse("admin_portal:grade_distribution_monitor_detail", args=[offering.id, self.period.id])
        self.assertEqual(self.client.get(url + "?" + response.context["page_query"]).status_code, 404)
        self.assertTrue(self.search(campus_id=campus.id).context["errors"])
        self.assertTrue(self.search(department_id=dept.id).context["errors"])

    def test_authorized_all_campuses_selected_campus_and_duplicate_department_names(self):
        campus, dept, offering = self._second_campus(grant_admin=True)
        self.client.force_login(self.admin_user)
        default = self.search()
        self.assertEqual(len(default.context["rows"]), 1)
        all_response = self.search(campus_id="all")
        self.assertEqual(len(all_response.context["rows"]), 2)
        labels = [d["label"] for d in all_response.context["filter_options"]["departments"]]
        self.assertEqual(len(labels), 2)
        self.assertEqual(len(set(labels)), 2)
        selected = self.search(campus_id=campus.id)
        self.assertEqual(selected.context["rows"][0]["course_code"], "CUB101")
        self.assertEqual([d["id"] for d in selected.context["filter_options"]["departments"]], [dept.id])
        self.assertEqual(self.detail(selected).status_code, 200)
        deny = UserPermission.objects.create(user=self.admin_user, permission=Permission.objects.get(code="grade_distribution_monitor.read"), tenant=self.tenant, campus=campus, grant_type="DENY")
        self.assertEqual(self.client.get(self.url + "?" + all_response.context["export_query"]).status_code, 400)
        narrowed = self.search(campus_id="all")
        self.assertEqual(len(narrowed.context["rows"]), 1)
        deny.delete()

    def test_foreign_tenant_offering_and_period_cannot_be_requested(self):
        other_tenant = Tenant.objects.create(code="OTHER", name="Other")
        other_template = GradingTemplate.objects.create(tenant=other_tenant, code="OTHER", name="Other")
        other_period = GradingTemplatePeriod.objects.create(template=other_template, code="PRELIM", name="Prelim")
        # Even inconsistent foreign physical-period links must fail closed.
        row = StudentPeriodGrade.objects.filter(offering=self.offering).first()
        row.template_period = other_period; row.save()
        self.client.force_login(self.admin_user)
        response = self.search()
        self.assertEqual(response.context["rows"][0]["graded_count"], 9)
        url = reverse("admin_portal:grade_distribution_monitor_detail", args=[self.offering.id, other_period.id])
        self.assertEqual(self.client.get(url + "?" + response.context["page_query"]).status_code, 404)
        self.assertTrue(self.search(academic_year_id=999999).context["errors"])

    def test_search_context_bound_to_user_and_topbar_scope(self):
        self.client.force_login(self.admin_user)
        response = self.search()
        other = self._user("other_review", "Other", "Review", self.tenant, self.campus, self.department)
        self._assign_permissions(other, include_monitor=True)
        self.client.force_login(other)
        self.assertEqual(self.detail(response).status_code, 400)

    def test_weighted_and_transmuted_gradebook_parity_uses_saved_values(self):
        self._formula_fixture(supplemental_count=5)
        TenantGradingProfile.objects.create(tenant=self.tenant, profile_code="DEPED", profile_name="DepEd", grading_template=self.template, period_grade_formula_mode="DEPED_TRANSMUTATION")
        ids = list(StudentPeriodGrade.objects.filter(offering=self.offering).values_list("student_id", flat=True))
        self._persist_formula_totals(ids)
        before = list(StudentPeriodGrade.objects.filter(offering=self.offering).values())
        data = self.detail(self.search()).json()
        for field, measure in (("period_grade", data), ("exam_grade", data["measures"][0]), ("class_standing_grade", data["measures"][1])):
            values = [r[field] for r in before if r[field] is not None]
            self.assertEqual(Decimal(measure["average"]), (sum(values) / len(values)).quantize(Decimal("0.01")))
        self.assertEqual(list(StudentPeriodGrade.objects.filter(offering=self.offering).values()), before)
        self.assertFalse(any("component grade" in m["unit"] for m in data["measures"]))

    def test_engine_attendance_interpretation_is_batched_and_read_only(self):
        from apps.attendance.models import AttendanceSession, AttendanceRecord
        c = GradingTemplateComponent.objects.create(template_period=self.period, code="CS", name="Attendance Standing", weight_percentage=100)
        GradingTemplateSubcomponent.objects.create(template_component=c, code="ATT", name="Attendance", weight_percentage=100, is_attendance_component=True)
        session = AttendanceSession.objects.create(tenant=self.tenant, campus=self.campus, offering=self.offering, template_period=self.period, session_date=date(2025, 7, 1), created_by_user=self.faculty_user)
        students = list(Student.objects.filter(enrollments__course_offering=self.offering).order_by("id")[:5])
        for student in students:
            AttendanceRecord.objects.create(session=session, student=student, status_code="P")
        self.client.force_login(self.admin_user)
        response = self.search()
        with CaptureQueriesContext(connection) as queries:
            data = self.detail(response).json()
        self.assertEqual([m["label"] for m in data["measures"]], ["Examination Grades", "Class Standing Grades"])
        self.assertEqual(data["measures"][1]["analytics_state"], "unavailable")
        self.assertEqual(data["incomplete_source_count"], 5)
        self.assertLessEqual(sum('attendance_records' in q["sql"] for q in queries), 2)

    def test_missing_assessment_routes_and_zero_sources_are_not_missing_saved_grades(self):
        c = GradingTemplateComponent.objects.create(template_period=self.period, code="EX", name="Exam", weight_percentage=100, is_exam_component=True)
        for code in ("A", "B"):
            GradingTemplateSubcomponent.objects.create(template_component=c, code=code, name=code, weight_percentage=50)
        sub = c.subcomponents.get(code="A")
        activity = GradeActivity.objects.create(tenant=self.tenant, campus=self.campus, offering=self.offering, template_period=self.period, template_component=c, template_subcomponent=sub, title="Zero exam", total_score=100, created_by_user=self.faculty_user)
        students = list(Student.objects.filter(enrollments__course_offering=self.offering).order_by("id")[:5])
        for student in students:
            StudentActivityScore.objects.create(activity=activity, student=student, raw_score=0, computed_score=50)
        self.client.force_login(self.admin_user)
        data = self.detail(self.search()).json()
        self.assertEqual(data["missing_count"], 0)
        self.assertEqual(data["incomplete_source_count"], 10)
        self.assertFalse(any("Actual examination percentage" in m["label"] for m in data["measures"]))
        configured = next(m for m in data["measures"] if m["label"] == "Zero exam")
        self.assertEqual(configured["graded_count"], 5)
        self.assertEqual(Decimal(configured["minimum"]), 50)
        self.assertEqual(set(StudentActivityScore.objects.filter(activity=activity).values_list("raw_score", flat=True)), {Decimal("0")})
        self.assertEqual(Decimal(data["average"]), Decimal("90.40"))

    def test_component_activity_limit_is_explicit_and_saved_totals_remain_available(self):
        c = GradingTemplateComponent.objects.create(template_period=self.period, code="EX", name="Exam", weight_percentage=100, is_exam_component=True)
        GradeActivity.objects.bulk_create([GradeActivity(tenant=self.tenant, campus=self.campus, offering=self.offering, template_period=self.period, template_component=c, title=f"Exam {i}", total_score=100, created_by_user=self.faculty_user) for i in range(201)])
        self.client.force_login(self.admin_user)
        data = self.detail(self.search()).json()
        self.assertEqual(len(data["measures"]), 2)
        self.assertIn("200 activities", " ".join(data["notes"]))
        self.assertIsNone(data["incomplete_source_count"])
        self.assertEqual(data["graded_count"], 10)
        self._privacy_assert_payload(data)
        self._privacy_assert_payload(self.search().context["rows"][0])

    def test_superuser_explicit_monitor_deny_has_no_dataset_access(self):
        self.admin_user.is_superuser = True; self.admin_user.save()
        UserPermission.objects.create(user=self.admin_user, permission=Permission.objects.get(code="grade_distribution_monitor.read"), tenant=self.tenant, campus=self.campus, grant_type="DENY")
        self.client.force_login(self.admin_user)
        response = self.search()
        self.assertEqual(response.context["rows"], [])
        self.assertTrue(response.context["errors"])

    def test_missing_historical_roster_explains_omitted_saved_records(self):
        self.offering.status = "ARCHIVED"; self.offering.save()
        Enrollment.objects.filter(course_offering=self.offering).delete()
        self.client.force_login(self.admin_user)
        data = self.detail(self.search()).json()
        self.assertEqual(data["eligible_count"], 0)
        self.assertEqual(data["graded_count"], 0)
        self.assertIn("10 saved record(s)", " ".join(data["notes"]))
        self.assertIn("No dependable eligible enrollment", " ".join(data["notes"]))

    def test_ended_open_term_uses_retained_historical_assignments_and_roster(self):
        self.term.end_date = date(2025, 6, 30)
        self.term.save()
        FacultyAssignment.objects.filter(offering=self.offering).update(is_active=False)
        Enrollment.objects.filter(course_offering=self.offering).update(is_active=False)
        Student.objects.filter(tenant=self.tenant).update(is_active=False)
        self.client.force_login(self.admin_user)
        response = self.search()
        row = response.context["rows"][0]
        self.assertTrue(row["historical"])
        self.assertEqual(row["graded_count"], 10)
        self.assertEqual(row["average"], Decimal("90.40"))
        self.assertEqual(self.detail(response).status_code, 200)

    def test_faculty_department_filters_preserve_default_scope_role_fallbacks(self):
        other_department = Department.objects.create(tenant=self.tenant, campus=self.campus, code="DEFAULT", name="Default Department")
        self.faculty_user.default_department = other_department
        self.faculty_user.save()
        UserRole.objects.filter(user=self.faculty_user, role__code="FACULTY").update(tenant=None, campus=None)
        self.client.force_login(self.admin_user)
        response = self.search(department_id=self.department.id)
        self.assertEqual([d["id"] for d in response.context["filter_options"]["departments"]], [self.department.id])
        self.assertEqual(len(response.context["rows"]), 1)
        self.assertEqual(self.detail(response).status_code, 200)
        self.assertEqual(self.client.get(self.url + "?" + response.context["export_query"]).status_code, 200)

    def _user(self, username, first_name, last_name, tenant, campus, department):
        return User.objects.create_user(
            username=username,
            email=f"{username}@example.com",
            password="testpass123",
            first_name=first_name,
            last_name=last_name,
            default_tenant=tenant,
            default_campus=campus,
            default_department=department,
            privacy_consent_version=getattr(settings, "PRIVACY_CONSENT_VERSION", "2026-03"),
            privacy_consent_at=timezone.now(),
        )

    def _assign_permissions(self, user, *, include_monitor):
        role, _ = Role.objects.get_or_create(code=f"ROLE_{user.username}", defaults={"name": f"Role {user.username}"})
        admin_access, _ = Permission.objects.get_or_create(
            code="admin_portal.access",
            defaults={"module": "admin_portal", "action": "access"},
        )
        RolePermission.objects.get_or_create(role=role, permission=admin_access)
        if include_monitor:
            monitor, _ = Permission.objects.get_or_create(
                code="grade_distribution_monitor.read",
                defaults={"module": "grade_distribution_monitor", "action": "read"},
            )
            RolePermission.objects.get_or_create(role=role, permission=monitor)
        UserRole.objects.create(
            user=user,
            role=role,
            tenant=user.default_tenant,
            campus=user.default_campus,
            department=user.default_department,
        )

    def _assign_faculty_role(self, user):
        role, _ = Role.objects.get_or_create(code="FACULTY", defaults={"name": "Faculty"})
        UserRole.objects.create(
            user=user,
            role=role,
            tenant=user.default_tenant,
            campus=user.default_campus,
            department=user.default_department,
        )

    def _seed_grades(self, offering, values, period=None, extra_enrolled_without_grade=0):
        period = period or self.period
        for index, value in enumerate(values, start=1):
            student = Student.objects.create(
                tenant=offering.tenant,
                campus=offering.campus,
                department=offering.department,
                program=offering.program,
                student_no=f"{offering.course.code}-{index:03d}",
                last_name=f"Student{index}",
                first_name="Test",
            )
            Enrollment.objects.create(
                tenant=offering.tenant,
                campus=offering.campus,
                academic_year=offering.academic_year,
                term=offering.term,
                student=student,
                course_offering=offering,
            )
            StudentPeriodGrade.objects.create(
                tenant=offering.tenant,
                campus=offering.campus,
                offering=offering,
                template_period=period,
                student=student,
                period_grade=Decimal(str(value)),
            )
        for index in range(len(values) + 1, len(values) + extra_enrolled_without_grade + 1):
            student = Student.objects.create(
                tenant=offering.tenant,
                campus=offering.campus,
                department=offering.department,
                program=offering.program,
                student_no=f"{offering.course.code}-{index:03d}",
                last_name=f"Student{index}",
                first_name="NoGrade",
            )
            Enrollment.objects.create(
                tenant=offering.tenant,
                campus=offering.campus,
                academic_year=offering.academic_year,
                term=offering.term,
                student=student,
                course_offering=offering,
            )

    def _create_extra_offering(self, course_code, course_title, section_code):
        course = Course.objects.create(
            tenant=self.tenant,
            campus=self.campus,
            department=self.department,
            code=course_code,
            title=course_title,
        )
        section = Section.objects.create(
            tenant=self.tenant,
            campus=self.campus,
            department=self.department,
            program=self.program,
            code=section_code,
            name=section_code,
        )
        offering = CourseOffering.objects.create(
            tenant=self.tenant,
            campus=self.campus,
            department=self.department,
            program=self.program,
            academic_year=self.academic_year,
            term=self.term,
            course=course,
            section=section,
        )
        FacultyAssignment.objects.create(
            tenant=self.tenant,
            campus=self.campus,
            offering=offering,
            faculty_user=self.faculty_user,
            response_status=FacultyAssignment.ResponseStatus.ACCEPTED,
            accepted_at=timezone.now(),
            is_primary=True,
        )
        return offering

    def _gradebook_activity_fixture(self, *, values=None, total=50, mode="RAW_BASE50", exam=False, title="Written assessment", students=None):
        component = GradingTemplateComponent.objects.create(template_period=self.period,
            code=f"ACT{GradeActivity.objects.count()}", name="Assessment", weight_percentage=100,
            is_exam_component=exam, score_input_mode=mode)
        activity = GradeActivity.objects.create(tenant=self.tenant, campus=self.campus, offering=self.offering,
            template_period=self.period, template_component=component, title=title, total_score=total)
        ids = students if students is not None else list(StudentPeriodGrade.objects.filter(offering=self.offering).order_by("student_id").values_list("student_id", flat=True))
        for sid, raw in zip(ids, values or [25] * len(ids)):
            computed = FacultyGradingService.compute_activity_score(raw_score=Decimal(str(raw)),
                total_score=Decimal(total), base_value=Decimal("50"), score_input_mode=mode)
            StudentActivityScore.objects.create(activity=activity, student_id=sid, raw_score=raw, computed_score=computed)
        self.client.force_login(self.admin_user)
        return activity, ids

    def test_base50_gradebook_activity_columns_match_analytics_frequencies(self):
        from collections import Counter
        from apps.grading.tabulation import PeriodSummaryLayoutService
        activity, ids = self._gradebook_activity_fixture(values=[0, 10, 20, 25, 30, 35, 40, 45, 49, 50])
        self._persist_formula_totals(ids)
        layout = PeriodSummaryLayoutService.build_layout(self.period, [activity])
        score_map = {(r.student_id, r.activity_id): r.computed_score for r in StudentActivityScore.objects.filter(activity=activity)}
        gradebook_values = []
        for student in Student.objects.filter(id__in=ids):
            row = PeriodSummaryLayoutService.build_row_values({"student": student}, layout, score_map)
            gradebook_values.extend(row["class_standing_blocks"][0]["sections"][0]["activity_values"])
        response = self.search()
        data = self.detail(response).json()
        measure = next(m for m in data["measures"] if m["label"] == activity.title)
        actual = {Decimal(f["grade"]): f["count"] for f in measure["frequency"]}
        self.assertEqual(actual, Counter(gradebook_values))
        self.assertIn(Decimal("75"), actual)  # Gradebook 25/50 -> Base 50 grade 75.
        self.assertNotIn(Decimal("25"), actual)
        self.assertEqual(Decimal(measure["average"]), sum(gradebook_values) / 10)
        for field, metric in (("period_grade", data), ("exam_grade", data["measures"][0]), ("class_standing_grade", data["measures"][1])):
            values = list(StudentPeriodGrade.objects.filter(offering=self.offering).exclude(**{field: None}).values_list(field, flat=True))
            self.assertEqual({Decimal(f["grade"]): f["count"] for f in metric["frequency"]}, Counter(values))
        workbook = self._workbook(self.client.get(self.url + "?" + response.context["export_query"]))
        records = self._worksheet_records(workbook, "Grade Distribution")
        self.assertEqual(len(records), 3)
        self.assertNotIn(activity.title, {r["Computed Grade Component"] for r in records})
        self._assert_count_export_reconciliation(response)

    def test_direct100_gradebook_ratings_use_exact_unrounded_breakdown(self):
        activity, ids = self._gradebook_activity_fixture(total=100, mode="DIRECT_PERCENTAGE", exam=True,
            values=[100, 99, 98, 97, 96, 95, Decimal("99.75"), 90, 89, 0])
        self._persist_formula_totals(ids)
        data = self.detail(self.search()).json()
        measure = next(m for m in data["measures"] if m["label"] == activity.title)
        self.assertEqual([r["count"] for r in measure["high_grade_breakdown"]], [1, 1, 1, 1, 1, 1, 6, 8])
        self.assertEqual(measure["exact_100_count"], 1)
        self.assertIn({"grade": "99.75", "count": 1, "percentage": "10.00"}, measure["frequency"])

    def test_monitor_reads_saved_configured_edits_without_recomputing(self):
        activity, ids = self._gradebook_activity_fixture()
        StudentActivityScore.objects.filter(activity=activity).update(computed_score=83)
        self._persist_formula_totals(ids)
        before_scores = list(StudentActivityScore.objects.values())
        before_grades = list(StudentPeriodGrade.objects.values())
        with patch.object(FacultyGradingService, "recompute_period_summary_for_students", side_effect=AssertionError("write forbidden")):
            data = self.detail(self.search()).json()
        measure = next(m for m in data["measures"] if m["label"] == activity.title)
        self.assertEqual(Decimal(measure["average"]), 83)
        self.assertEqual(list(StudentActivityScore.objects.values()), before_scores)
        self.assertEqual(list(StudentPeriodGrade.objects.values()), before_grades)
        StudentActivityScore.objects.filter(activity=activity).update(is_active=False)
        measure = self.detail(self.search()).json()["measures"][-1]
        self.assertEqual(measure["analytics_state"], "unavailable")

    def test_historical_activity_grades_preserve_original_campus_eligibility(self):
        activity, ids = self._gradebook_activity_fixture()
        self._persist_formula_totals(ids)
        self.offering.status = "ARCHIVED"
        self.offering.save()
        response = self.search()
        before = self.detail(response).json()
        target_campus, department, target = self._second_campus()
        Student.objects.filter(pk__in=ids).update(campus=target_campus, department=department, program=target.program, is_active=False)
        after = self.detail(self.search()).json()
        self.assertEqual(before["measures"], after["measures"])
        self.assertEqual(Decimal(after["measures"][-1]["average"]), 75)
        self.assertIn("score revisions are not snapshotted", " ".join(after["notes"]))
        self.assertIsNone(after["incomplete_source_count"])
        StudentActivityScore.objects.filter(activity=activity).update(computed_score=84)
        after = self.detail(self.search()).json()
        self.assertEqual(Decimal(after["measures"][-1]["average"]), 84)

    def test_historical_sparse_activity_blocks_saved_formula_ancestors(self):
        activity, ids = self._gradebook_activity_fixture(students=list(StudentPeriodGrade.objects.filter(offering=self.offering).values_list("student_id", flat=True))[:1])
        StudentPeriodGrade.objects.filter(offering=self.offering).update(exam_grade=75, class_standing_grade=80)
        self.offering.status = "ARCHIVED"
        self.offering.save()
        response = self.search()
        data = self.detail(response).json()
        for metric in [data, *data["measures"]]:
            self._privacy_assert_payload(metric)
        export = self._export_text(self.client.get(self.url + "?" + response.context["export_query"]))
        self.assertNotIn("75.0", export)
        self.assertNotIn("80.0", export)

    def test_overlapping_activity_populations_cannot_isolate_one_student(self):
        ids = list(StudentPeriodGrade.objects.filter(offering=self.offering).order_by("student_id").values_list("student_id", flat=True))
        first, _ = self._gradebook_activity_fixture(students=ids[:5], title="First assessment")
        second, _ = self._gradebook_activity_fixture(students=ids[4:9], title="Second assessment")
        data = self.detail(self.search()).json()
        measures = {m["label"]: m for m in data["measures"]}
        self.assertEqual(measures[first.title]["analytics_state"], "disclosed")
        self._privacy_assert_payload(measures[second.title])
        self.assertEqual(measures[second.title]["privacy_reason"], "complementary_population")
        response = self.search()
        workbook = self._workbook(self.client.get(self.url + "?" + response.context["export_query"]))
        self.assertEqual(workbook.sheetnames, ["Grade Distribution"])
        self._assert_count_export_reconciliation(response)

    def test_activity_status_complements_and_five_student_boundary(self):
        activity, ids = self._gradebook_activity_fixture()
        for count in (1, 4, 5, 6, 9, 10):
            with self.subTest(finalized=count):
                StudentPeriodGrade.objects.filter(offering=self.offering).update(is_finalized=False)
                StudentPeriodGrade.objects.filter(student_id__in=ids[:count], offering=self.offering).update(is_finalized=True)
                for status in ("all", "finalized", "unfinalized"):
                    response = self.search(status=status)
                    metric = self.detail(response).json()["measures"][-1]
                    chosen = 10 if status == "all" else count if status == "finalized" else 10-count
                    state = "unavailable" if chosen == 0 else "disclosed" if count in (5, 10) else "privacy_suppressed"
                    self.assertEqual(metric["analytics_state"], state)
                    if state == "privacy_suppressed":
                        self._privacy_assert_payload(metric)

    def test_xlsx_styles_types_filters_suppression_and_literal_user_text(self):
        activity, ids = self._gradebook_activity_fixture()
        self._persist_formula_totals(ids)
        self.course.title = '=HYPERLINK("https://invalid.example","x")'
        self.course.save()
        self.faculty_user.last_name = '+Literal'
        self.faculty_user.save()
        response = self.search()
        exported = self.client.get(self.url + "?" + response.context["export_query"])
        workbook = self._workbook(exported)
        self.assertEqual(workbook.sheetnames, ["Grade Distribution"])
        self.assertIn('.xlsx"', exported["Content-Disposition"])
        sheet = workbook.active
        self.assertEqual(sheet.freeze_panes, "F2")
        self.assertEqual(sheet.auto_filter.ref, sheet.dimensions)
        self.assertTrue(sheet['A1'].font.bold)
        self.assertEqual(sheet['A1'].fill.fgColor.rgb, '0020613C')
        self.assertGreaterEqual(sheet.column_dimensions['A'].width, 30)
        self.assertFalse(any(c.data_type == 'f' for cells in sheet for c in cells))
        self.assertEqual(sheet['A2'].data_type, 's')
        self.assertTrue(sheet['A2'].value.startswith('+Literal'))
        self.assertEqual(sheet['C2'].data_type, 's')
        self.assertIn('=HYPERLINK', sheet['C2'].value)
        self.assertEqual(sheet['F3'].number_format, '0')
        StudentActivityScore.objects.filter(activity=activity).exclude(student_id__in=ids[:4]).update(is_active=False)
        self._assert_count_export_reconciliation(self.search())

    def test_xlsx_aggregate_row_limit_is_exact_and_does_not_truncate(self):
        self.client.force_login(self.admin_user)
        response = self.search()
        url = self.url + '?' + response.context['export_query']
        with patch('apps.admin_portal.grade_distribution_export.MAX_WORKBOOK_ROWS', 3):
            self.assertEqual(self.client.get(url).status_code, 200)
        with patch('apps.admin_portal.grade_distribution_export.MAX_WORKBOOK_ROWS', 2):
            rejected = self.client.get(url)
            self.assertEqual(rejected.status_code, 400)
            self.assertIn('no records were truncated', ' '.join(rejected.json()['errors']))
        # The old download URL remains usable but returns the new workbook.
        self._workbook(self.client.get(url.replace('export=xlsx', 'export=csv')))

    def test_xlsx_row_limit_stops_later_class_source_queries(self):
        from apps.admin_portal.grade_distribution_metrics import component_measures
        other = self._create_extra_offering("NEXT", "Later class", "NEXT-A")
        self._seed_grades(other, [90] * 10)
        self.client.force_login(self.admin_user)
        response = self.search()
        with patch("apps.admin_portal.grade_distribution.component_measures", wraps=component_measures) as source:
            with patch("apps.admin_portal.grade_distribution_export.MAX_WORKBOOK_ROWS", 2):
                rejected = self.client.get(self.url + "?" + response.context["export_query"])
        self.assertEqual(rejected.status_code, 400)
        self.assertEqual(source.call_count, 0)
        self.assertIn("no records were truncated", " ".join(rejected.json()["errors"]))

    def test_inactive_activity_route_fails_closed_without_invented_classification(self):
        activity, ids = self._gradebook_activity_fixture()
        activity.template_component.is_active = False
        activity.template_component.save()
        data = self.detail(self.search()).json()
        self._privacy_assert_payload(data)
        self._privacy_assert_payload(data['measures'][-1])
        self.assertNotIn('subjective', data['measures'][-1])
        self.assertIn('inactive grading route', ' '.join(data['notes']))

    def test_real_formula_payload_for_frontend_dom(self):
        import os
        from pathlib import Path
        self._formula_fixture()
        data = self.detail(self.search()).json()
        self._privacy_assert_payload(data)
        self.assertTrue(any(m["analytics_state"] == "disclosed" for m in data["measures"]))
        if os.environ.get("TMP_MONITOR_DOM_PAYLOAD"):
            Path(os.environ["TMP_MONITOR_DOM_PAYLOAD"]).write_text(json.dumps(data), encoding="utf-8")

    def _worksheet_records(self, workbook, sheet_name):
        rows = list(workbook[sheet_name].values)
        return [dict(zip(rows[0], values)) for values in rows[1:]]

    def _assert_count_export_reconciliation(self, response):
        from apps.admin_portal.grade_distribution_export import HEADERS, export_components
        rows = response.context["rows"]
        workbook = self._workbook(self.client.get(self.url + "?" + response.context["export_query"]))
        self.assertEqual(workbook.sheetnames, ["Grade Distribution"])
        self.assertEqual([c.value for c in workbook.active[1]], HEADERS)
        records = self._worksheet_records(workbook, "Grade Distribution")
        self.assertEqual(len(records), 3 * len(rows))
        for row_index, row in enumerate(rows):
            for offset, (label, metric) in enumerate(export_components(row)):
                record = records[3 * row_index + offset]
                self.assertEqual(record["Computed Grade Component"], label)
                expected = [m["count"] for m in metric["high_grade_breakdown"][:6]] if metric["analytics_state"] == "disclosed" else [None] * 6
                self.assertEqual([record[f"Exactly {g}"] for g in range(100, 94, -1)], expected)
        return records

    def _duplicate_quiz_fixture(self):
        component = GradingTemplateComponent.objects.create(template_period=self.period,
            code="QUIZZES", name="Class Standing", weight_percentage=100)
        activities = []
        ids = list(StudentPeriodGrade.objects.filter(offering=self.offering).order_by("student_id").values_list("student_id", flat=True))
        for code, name, raw in (("WRITTEN", "Written Quizzes", 25), ("ORAL", "Oral Quizzes", 40)):
            sub = GradingTemplateSubcomponent.objects.create(template_component=component, code=code, name=name, weight_percentage=50)
            activity = GradeActivity.objects.create(tenant=self.tenant, campus=self.campus, offering=self.offering,
                template_period=self.period, template_component=component, template_subcomponent=sub, title="Quiz", total_score=50)
            grade = FacultyGradingService.compute_activity_score(raw_score=Decimal(raw), total_score=Decimal(50), base_value=Decimal(50))
            for sid in ids:
                StudentActivityScore.objects.create(activity=activity, student_id=sid, raw_score=raw, computed_score=grade)
            activities.append(activity)
        self._persist_formula_totals(ids)
        self.client.force_login(self.admin_user)
        return activities, ids

    def test_dataset_identity_preserves_two_physical_prelim_periods_in_xlsx(self):
        template = GradingTemplate.objects.create(tenant=self.tenant, code="HISTORICAL", name="Historical Template")
        second = GradingTemplatePeriod.objects.create(template=template, code="PRELIM", name="Prelim", sequence_no=1)
        self.offering.status = "ARCHIVED"
        self.offering.save()
        existing = list(StudentPeriodGrade.objects.filter(offering=self.offering))
        StudentPeriodGrade.objects.filter(offering=self.offering).update(period_grade=75)
        StudentPeriodGrade.objects.bulk_create([StudentPeriodGrade(tenant=self.tenant, campus=self.campus,
            offering=self.offering, template_period=second, student_id=r.student_id, period_grade=90) for r in existing])
        self.client.force_login(self.admin_user)
        response = self.search()
        rows = response.context["rows"]
        self.assertEqual(len(rows), 2)
        self.assertEqual({r["period"] for r in rows}, {"Prelim"})
        self.assertEqual({r["physical_period_id"] for r in rows}, {self.period.id, second.id})
        self.assertEqual(len({r["dataset_id"] for r in rows}), 2)
        for row in rows:
            self.assertEqual(self.detail(response, row).json()["dataset_id"], row["dataset_id"])
            self.assertContains(response, f'Period #{row["physical_period_id"]}')
        import os
        from pathlib import Path
        if os.environ.get("TMP_PERIOD_DATASET_DOM_PAYLOAD"):
            Path(os.environ["TMP_PERIOD_DATASET_DOM_PAYLOAD"]).write_text(
                json.dumps([self.detail(response, row).json() for row in rows]), encoding="utf-8")
        records = self._assert_count_export_reconciliation(response)
        self.assertEqual(len({r["Grading Period"] for r in records}), 2)
        self.assertTrue(all("Prelim" in r["Grading Period"] for r in records))

    def test_duplicate_quiz_routes_reconcile_and_match_gradebook_for_all_statuses(self):
        from apps.grading.tabulation import PeriodSummaryLayoutService
        import os
        from pathlib import Path
        activities, ids = self._duplicate_quiz_fixture()
        StudentPeriodGrade.objects.filter(offering=self.offering, student_id__in=ids[:5]).update(is_finalized=True)
        layout = PeriodSummaryLayoutService.build_layout(self.period, activities)
        score_map = {(r.student_id, r.activity_id): r.computed_score for r in StudentActivityScore.objects.filter(activity__in=activities)}
        source_ids = None
        selection_ids = set()
        for status, count in (("all", 10), ("finalized", 5), ("unfinalized", 5)):
            with self.subTest(status=status):
                response = self.search(status=status)
                data = self.detail(response).json()
                quizzes = [m for m in data["measures"] if m["label"] == "Quiz"]
                self.assertEqual(len(quizzes), 2)
                self.assertEqual({m["activity_id"] for m in quizzes}, {a.id for a in activities})
                self.assertEqual(len({m["dataset_id"] for m in quizzes}), 2)
                current_sources = {m["dataset_id"].split("/status:")[0] for m in quizzes}
                self.assertEqual(current_sources, source_ids or current_sources)
                source_ids = current_sources
                self.assertTrue(selection_ids.isdisjoint(m["dataset_id"] for m in quizzes))
                selection_ids.update(m["dataset_id"] for m in quizzes)
                for metric in quizzes:
                    self.assertEqual(metric["graded_count"], count)
                    self.assertEqual(metric["selection_status"], status)
                    self.assertIn(f'Activity #{metric["activity_id"]}', metric["gradebook_context"])
                    self.assertIn("Quizzes", metric["gradebook_context"])
                sample = Student.objects.get(pk=ids[0])
                gradebook = PeriodSummaryLayoutService.build_row_values({"student": sample}, layout, score_map)
                gradebook_by_activity = {
                    aid: value
                    for section, values in zip(layout["class_standing_blocks"][0]["sections"], gradebook["class_standing_blocks"][0]["sections"])
                    for aid, value in zip(section["activity_ids"], values["activity_values"])
                }
                self.assertEqual(gradebook_by_activity,
                                 {m["activity_id"]: Decimal(m["frequency"][0]["grade"]) for m in quizzes})
                records = self._assert_count_export_reconciliation(response)
                self.assertEqual(data["quiz"]["graded_count"], count)
                self.assertEqual(len([r for r in records if r["Computed Grade Component"] == "Quizzes"]), 1)
                if status == "all" and os.environ.get("TMP_DATASET_DOM_PAYLOAD"):
                    Path(os.environ["TMP_DATASET_DOM_PAYLOAD"]).write_text(json.dumps(data), encoding="utf-8")

    def test_dataset_identity_survives_title_edits_and_historical_retention(self):
        activities, ids = self._duplicate_quiz_fixture()
        response = self.search()
        before = self.detail(response).json()
        before_ids = {m["activity_id"]: m["dataset_id"] for m in before["measures"] if m["activity_id"]}
        GradeActivity.objects.filter(pk=activities[0].pk).update(title="Quiz revised")
        self.offering.status = "ARCHIVED"
        self.offering.save()
        after = self.detail(self.search()).json()
        after_ids = {m["activity_id"]: m["dataset_id"] for m in after["measures"] if m["activity_id"]}
        self.assertEqual(before_ids, after_ids)
        self.assertIn("Quiz revised", {m["label"] for m in after["measures"]})
        self._assert_count_export_reconciliation(self.search())

    def test_equal_quiz_maps_reuse_exact_dataset_reference_but_keep_export_records(self):
        activities, ids = self._duplicate_quiz_fixture()
        # Keep saved buckets at 82.50 so the two equal Quiz maps refer to each
        # other, rather than to a same-valued saved period/standing dataset.
        StudentActivityScore.objects.filter(activity=activities[1]).update(computed_score=75)
        response = self.search()
        data = self.detail(response).json()
        quizzes = [m for m in data["measures"] if m["label"] == "Quiz"]
        self.assertNotEqual(quizzes[0]["dataset_id"], quizzes[1]["dataset_id"])
        self.assertEqual(quizzes[1]["equivalent_to"], "Quiz")
        self.assertEqual(quizzes[1]["equivalent_dataset_id"], quizzes[0]["dataset_id"])
        self.assertEqual(quizzes[1]["equivalent_context"], quizzes[0]["gradebook_context"])
        self.assertIn("Written Quizzes", quizzes[1]["equivalent_context"])
        records = self._assert_count_export_reconciliation(response)
        self.assertEqual(len(records), 3)
        self.assertEqual(data["quiz"]["analytics_state"], "unavailable")

    def test_export_limit_counts_only_three_components_per_physical_record(self):
        self._duplicate_quiz_fixture()
        response = self.search()
        url = self.url + "?" + response.context["export_query"]
        with patch("apps.admin_portal.grade_distribution_export.MAX_WORKBOOK_ROWS", 3):
            self.assertEqual(self.client.get(url).status_code, 200)
        with patch("apps.admin_portal.grade_distribution_export.MAX_WORKBOOK_ROWS", 2):
            self.assertEqual(self.client.get(url).status_code, 400)

    def test_duplicate_quiz_identity_does_not_expose_formula_protected_xlsx_difference(self):
        ids, components = self._formula_fixture()
        GradeActivity.objects.filter(template_component__in=components[:2]).update(title="Quiz")
        response = self.search()
        data = self.detail(response).json()
        quizzes = [m for m in data["measures"] if m["label"] == "Quiz"]
        common = next(m for m in quizzes if m["graded_count"] == 10)
        sparse = next(m for m in quizzes if m["graded_count"] == 1)
        self.assertNotEqual(common["dataset_id"], sparse["dataset_id"])
        self._privacy_assert_payload(sparse)
        self._privacy_assert_payload(data)
        self._privacy_assert_payload(data["measures"][0])
        self.assertEqual(sum(StudentPeriodGrade.objects.filter(offering=self.offering).values_list("exam_grade", flat=True)), 788)
        records = self._assert_count_export_reconciliation(response)
        examination = next(r for r in records if r["Computed Grade Component"] == "Periodic Examination")
        self.assertTrue(all(examination[f"Exactly {g}"] is None for g in range(100, 94, -1)))
        self.assertNotIn("Average Grade", examination)
        self.assertNotIn("Grade frequencies", self._workbook(self.client.get(self.url + "?" + response.context["export_query"])).sheetnames)


    def _same_label_offering(self, *, department=None, program=None, separate_section=False, grade=99):
        department = department or self.department
        program = program or self.program
        section = self.section
        if separate_section:
            section = Section.objects.create(tenant=self.tenant, campus=self.campus, department=department,
                program=program, code=self.section.code, name=self.section.name)
        offering = CourseOffering.objects.create(tenant=self.tenant, campus=self.campus, department=department,
            program=program, academic_year=self.academic_year, term=self.term, course=self.course, section=section)
        FacultyAssignment.objects.create(tenant=self.tenant, campus=self.campus, offering=offering,
            faculty_user=self.faculty_user, response_status="ACCEPTED", accepted_at=timezone.now())
        for row in StudentPeriodGrade.objects.filter(offering=self.offering):
            Enrollment.objects.create(tenant=self.tenant, campus=self.campus, academic_year=self.academic_year,
                term=self.term, course_offering=offering, student_id=row.student_id)
            StudentPeriodGrade.objects.create(tenant=self.tenant, campus=self.campus, offering=offering,
                template_period=self.period, student_id=row.student_id, period_grade=grade,
                exam_grade=grade, class_standing_grade=grade)
        StudentPeriodGrade.objects.filter(offering=self.offering).update(period_grade=100, exam_grade=100, class_standing_grade=100)
        self.client.force_login(self.admin_user)
        return offering

    def _class_label_export(self, response, count):
        workbook = self._workbook(self.client.get(self.url + "?" + response.context["export_query"]))
        self.assertEqual(workbook.sheetnames, ["Grade Distribution"])
        self.assertEqual(workbook.active.max_column, 11)
        records = list(workbook.active.values)[1:]
        self.assertEqual(len(records), 3 * count)
        self.assertEqual(len({tuple(row[:5]) for row in records}), 3 * count)
        return records

    def test_excel_class_labels_distinguish_exact_cross_department_collision(self):
        department = Department.objects.create(tenant=self.tenant, campus=self.campus, code="BBA", name="Business")
        program = Program.objects.create(tenant=self.tenant, campus=self.campus, department=department,
            code="BBA", name="Business Administration")
        other = self._same_label_offering(department=department, program=program)
        response = self.search()
        self.assertEqual({r["offering_id"] for r in response.context["rows"]}, {self.offering.id, other.id})
        self.assertEqual(len({(r["faculty_name"], r["campus"], r["course_code"], r["section"], r["period"])
                              for r in response.context["rows"]}), 1)
        records = self._class_label_export(response, 2)
        base = f"{self.course.code} - {self.course.title} / {self.section.code}"
        self.assertEqual({r[2] for r in records}, {base + " [Department: BSA]", base + " [Department: BBA]"})
        self.assertEqual({r[3] for r in records}, {"Prelim"})
        examination = {r[2]: r[5:] for r in records if r[4] == "Periodic Examination"}
        self.assertEqual(examination[base + " [Department: BSA]"], (10, 0, 0, 0, 0, 0))
        self.assertEqual(examination[base + " [Department: BBA]"], (0, 10, 0, 0, 0, 0))

    def test_excel_class_labels_use_program_only_when_department_cannot_distinguish(self):
        program = Program.objects.create(tenant=self.tenant, campus=self.campus, department=self.department,
            code="BBA", name="Business Administration")
        self._same_label_offering(program=program, separate_section=True)
        records = self._class_label_export(self.search(), 2)
        self.assertTrue(all("[Department:" not in r[2] for r in records))
        self.assertEqual({r[2].split(" [Program: ")[1] for r in records}, {"BSA]", "BBA]"})

    def test_excel_class_labels_qualify_only_remaining_program_collision(self):
        program = Program.objects.create(tenant=self.tenant, campus=self.campus, department=self.department,
            code="BBA", name="Business Administration")
        self._same_label_offering(program=program, separate_section=True)
        department = Department.objects.create(tenant=self.tenant, campus=self.campus, code="EDUC", name="Education")
        education = Program.objects.create(tenant=self.tenant, campus=self.campus, department=department,
            code="EDUC", name="Education")
        self._same_label_offering(department=department, program=education, grade=98)
        records = self._class_label_export(self.search(), 3)
        labels = {r[2] for r in records}
        self.assertEqual(len(labels), 3)
        self.assertTrue(any(label.endswith("[Department: EDUC]") for label in labels))
        self.assertTrue(any(label.endswith("[Department: BSA] [Program: BSA]") for label in labels))
        self.assertTrue(any(label.endswith("[Department: BSA] [Program: BBA]") for label in labels))

    def test_excel_class_labels_handle_missing_offering_program(self):
        program = Program.objects.create(tenant=self.tenant, campus=self.campus, department=self.department,
            code="BBA", name="Business Administration")
        other = self._same_label_offering(program=program, separate_section=True)
        CourseOffering.objects.filter(pk=other.pk).update(program=None)
        records = self._class_label_export(self.search(), 2)
        self.assertTrue(any(r[2].endswith("[Program: Not specified]") for r in records))
        self.assertTrue(any(r[2].endswith("[Program: BSA]") for r in records))

    def test_excel_class_labels_reject_remaining_ambiguous_physical_classes(self):
        program = Program.objects.create(tenant=self.tenant, campus=self.campus, department=self.department,
            code="BBA", name="Business Administration")
        other = self._same_label_offering(program=program, separate_section=True)
        # Different physical sections, but even the offering qualifiers agree.
        # Do not invent an identifier from row order or silently combine grades.
        CourseOffering.objects.filter(pk=other.pk).update(program=self.program)
        response = self.search()
        self.assertEqual(len(response.context["rows"]), 2)
        exported = self.client.get(self.url + "?" + response.context["export_query"])
        self.assertEqual(exported.status_code, 400)
        self.assertIn("class labels remain ambiguous", " ".join(exported.json()["errors"]))
        self.assertFalse(exported.has_header("Content-Disposition"))
        self.assertEqual(StudentPeriodGrade.objects.filter(offering_id__in=[self.offering.id, other.id]).count(), 20)

    def test_excel_class_labels_reject_missing_program_marker_collision(self):
        program = Program.objects.create(tenant=self.tenant, campus=self.campus, department=self.department,
            code="BBA", name="Business Administration")
        other = self._same_label_offering(program=program, separate_section=True)
        Program.objects.filter(pk=self.program.pk).update(code="Not specified")
        CourseOffering.objects.filter(pk=other.pk).update(program=None)
        response = self.search()
        exported = self.client.get(self.url + "?" + response.context["export_query"])
        self.assertEqual(exported.status_code, 400)
        self.assertIn("class labels remain ambiguous", " ".join(exported.json()["errors"]))
        self.assertFalse(exported.has_header("Content-Disposition"))

    def test_excel_class_labels_preserve_suppression_for_cross_department_collision(self):
        department = Department.objects.create(tenant=self.tenant, campus=self.campus, code="BBA", name="Business")
        other = self._same_label_offering(department=department)
        saved = StudentPeriodGrade.objects.filter(offering=other)
        keep = list(saved.order_by("id").values_list("id", flat=True)[:4])
        saved.exclude(pk__in=keep).update(period_grade=None, exam_grade=None, class_standing_grade=None)
        response = self.search()
        records = self._class_label_export(response, 2)
        private_rows = [r for r in records if "[Department: BBA]" in r[2]]
        self.assertEqual(len(private_rows), 3)
        self.assertTrue(all(value is None for row in private_rows for value in row[5:]))
        public = next(r for r in records if "[Department: BSA]" in r[2] and r[4] == "Periodic Examination")
        self.assertEqual(public[5], 10)

    def test_excel_class_labels_preserve_classes_and_physical_periods_together(self):
        department = Department.objects.create(tenant=self.tenant, campus=self.campus, code="BBA", name="Business")
        other = self._same_label_offering(department=department)
        template = GradingTemplate.objects.create(tenant=self.tenant, code="LEGACY-LABEL", name="Legacy Template")
        period = GradingTemplatePeriod.objects.create(template=template, code="PRELIM", name="Prelim", sequence_no=1)
        CourseOffering.objects.filter(pk__in=[self.offering.pk, other.pk]).update(status="ARCHIVED")
        for row in list(StudentPeriodGrade.objects.filter(offering_id__in=[self.offering.pk, other.pk])):
            row.pk = None
            row.template_period = period
            row.save(force_insert=True)
        response = self.search()
        records = self._class_label_export(response, 4)
        self.assertEqual(len({r[2] for r in records}), 2)
        self.assertEqual(len({r[3] for r in records}), 2)
        self.assertTrue(all("[Program:" not in r[2] for r in records))
        self.assertTrue(all("Legacy Template" in r[3] or "General Education" in r[3] for r in records))

    def test_excel_class_labels_reject_case_and_whitespace_equivalent_identifiers(self):
        from django.core.exceptions import ValidationError
        from apps.admin_portal.grade_distribution_export import class_labels
        row = dict(faculty_name="Faculty", campus="FV", course_code="C", course_title="Course", section="S",
            period="Prelim", physical_period_id=1, department="BSA", program="BSA")
        with self.assertRaisesMessage(ValidationError, "class labels remain ambiguous"):
            class_labels([dict(row, offering_id=1), dict(row, offering_id=2, department=" bsa ", program="bsa")])

    def test_excel_class_labels_leave_unambiguous_class_unqualified(self):
        self.client.force_login(self.admin_user)
        records = self._class_label_export(self.search(), 1)
        self.assertEqual({r[2] for r in records}, {f"{self.course.code} - {self.course.title} / {self.section.code}"})
        self.assertEqual({r[3] for r in records}, {"Prelim"})

    def _calculated_quiz_fixture(self, values=None):
        values = values or [Decimal("100"), Decimal("99.75"), Decimal("99"), Decimal("98"),
                            Decimal("97"), Decimal("96"), Decimal("95"), Decimal("94.5"), Decimal("50"), Decimal("90")]
        ids = list(StudentPeriodGrade.objects.filter(offering=self.offering).order_by("student_id").values_list("student_id", flat=True))
        standing = GradingTemplateComponent.objects.create(template_period=self.period, code="CS", name="Class Standing", weight_percentage=60)
        sub = GradingTemplateSubcomponent.objects.create(template_component=standing, code="QUIZZES", name="Quizzes", weight_percentage=100)
        activities = []
        for title in ("Q1", "Q2"):
            activity = GradeActivity.objects.create(tenant=self.tenant, campus=self.campus, offering=self.offering,
                template_period=self.period, template_component=standing, template_subcomponent=sub, title=title, total_score=50)
            for sid, value in zip(ids, values):
                StudentActivityScore.objects.create(activity=activity, student_id=sid, raw_score=value - 50, computed_score=value)
            activities.append(activity)
        exam = GradingTemplateComponent.objects.create(template_period=self.period, code="EXAM", name="Periodic Examination", weight_percentage=40, is_exam_component=True, score_input_mode="DIRECT_PERCENTAGE")
        exam_activity = GradeActivity.objects.create(tenant=self.tenant, campus=self.campus, offering=self.offering,
            template_period=self.period, template_component=exam, title="Exam", total_score=100)
        for sid in ids:
            StudentActivityScore.objects.create(activity=exam_activity, student_id=sid, raw_score=96, computed_score=96)
        self._persist_formula_totals(ids)
        self.client.force_login(self.admin_user)
        return ids, activities

    def test_default_sort_is_numeric_global_stable_and_export_preserves_filters(self):
        import time
        extra = []
        for index in range(26):
            offering = self._create_extra_offering(f"SORT{index:02}", "Sort course", f"S{index:02}")
            self._seed_grades(offering, [100] * (10 if index == 25 else 2) + [90] * (0 if index == 25 else 8))
            extra.append(offering)
        self.client.force_login(self.admin_user)
        start = time.perf_counter()
        with CaptureQueriesContext(connection) as queries:
            response = self.search(faculty_id=self.faculty_user.id)
        elapsed = time.perf_counter() - start
        query_count = len(queries)
        first = response.context["rows"]
        self.assertEqual(first[0]["offering_id"], extra[-1].id)
        self.assertEqual(first[0]["exact_100_count"], 10)
        self.assertEqual(first[1]["exact_100_count"], 2)
        self.assertEqual(response.context["selected"]["sort"], "exact_100_desc")
        second = self.client.get(self.url + "?" + response.context["page_query"] + "&page=2")
        ids = [r["offering_id"] for r in first + second.context["rows"]]
        self.assertEqual(len(ids), 27)
        self.assertEqual(len(set(ids)), 27)
        self.assertEqual(len(first), 25)
        repeated = self.search(faculty_id=self.faculty_user.id)
        self.assertEqual([r["offering_id"] for r in repeated.context["rows"]], ids[:25])
        params = parse_qs(response.context["export_query"])
        self.assertEqual(params["faculty_id"], [str(self.faculty_user.id)])
        self.assertEqual(params["sort"], ["exact_100_desc"])
        sheet = self._workbook(self.client.get(self.url + "?" + response.context["export_query"])).active
        self.assertTrue(sheet["C2"].value.startswith("SORT25"))
        self.assertEqual(sheet.max_row, 82)
        self.assertGreater(query_count, 0)
        self.assertLess(query_count, 1500)
        print(f"GRADE100_PERFORMANCE: 27 classes / 270 students / page 1: {elapsed:.3f}s, {query_count} SQL queries")

    def test_sort_validation_context_binding_and_page_reset(self):
        self.client.force_login(self.admin_user)
        self.assertTrue(self.search(sort="raw_count").context["errors"])
        response = self.search(sort="faculty_asc")
        self.assertContains(response, 'name="sort"')
        self.assertNotContains(response, 'name="page"')
        url = self.url + "?" + response.context["export_query"]
        self.assertEqual(self.client.get(url.replace("sort=faculty_asc", "sort=exact_100_desc")).status_code, 400)
        self.assertEqual(self.client.get(url).status_code, 200)

    def test_suppressed_ranking_never_uses_hidden_counts(self):
        base = dict(faculty_name="Same", campus="FV", course_code="C", section="S", physical_period_id=1)
        rows = [dict(base, offering_id=2, analytics_state="privacy_suppressed", exact_100_count=999),
                dict(base, offering_id=1, analytics_state="privacy_suppressed", exact_100_count=1),
                dict(base, offering_id=3, analytics_state="disclosed", exact_100_count=0)]
        rank = lambda: [r["offering_id"] for r in sorted(rows, key=GradeDistributionMonitorService.sort_key)]
        self.assertEqual(rank(), [3, 1, 2])
        rows[0]["exact_100_count"], rows[1]["exact_100_count"] = 0, 2000
        self.assertEqual(rank(), [3, 1, 2])

    def test_calculated_quiz_exact_counts_match_gradebook_and_not_raw_values(self):
        from collections import Counter
        from apps.grading.tabulation import PeriodSummaryLayoutService
        ids, activities = self._calculated_quiz_fixture()
        response = self.search()
        quiz = self.detail(response).json()["quiz"]
        layout = PeriodSummaryLayoutService.build_layout(self.period, activities)
        score_map = {(r.student_id, r.activity_id): r.computed_score for r in StudentActivityScore.objects.filter(activity__in=activities)}
        expected = Counter(PeriodSummaryLayoutService.print_group_average(layout,
            PeriodSummaryLayoutService.build_row_values({"student": student}, layout, score_map), quizzes=True)
            for student in Student.objects.filter(id__in=ids))
        self.assertEqual({Decimal(f["grade"]): f["count"] for f in quiz["frequency"]}, expected)
        self.assertEqual(quiz["graded_count"], 10)
        self.assertEqual([b["count"] for b in quiz["high_grade_breakdown"][:6]], [1, 1, 1, 1, 1, 1])
        self.assertEqual(expected[Decimal("99.75")], 1)
        exported = self._privacy_xlsx_row(self.client.get(self.url + "?" + response.context["export_query"]), "Quizzes")
        self.assertEqual([exported[f"Exactly {g}"] for g in range(100, 94, -1)], ["1"] * 6)
        self.assertContains(response, "not make Quiz averages immutable finalized snapshots")
        self.assertEqual(self._privacy_xlsx_row(self.client.get(self.url + "?" + response.context["export_query"]))["Exactly 96"], "10")

    def test_quiz_counts_one_average_per_student_not_individual_quiz_frequencies(self):
        ids, activities = self._calculated_quiz_fixture([Decimal(100)] * 10)
        StudentActivityScore.objects.filter(activity=activities[1]).update(raw_score=48, computed_score=98)
        self._persist_formula_totals(ids)
        response = self.search()
        quiz = self.detail(response).json()["quiz"]
        self.assertEqual(quiz["frequency"], [{"grade": "99.00", "count": 10, "percentage": "100.00"}])
        exported = self._privacy_xlsx_row(self.client.get(self.url + "?" + response.context["export_query"]), "Quizzes")
        self.assertEqual((exported["Exactly 100"], exported["Exactly 99"]), ("0", "10"))

    def test_quiz_weighted_details_and_average_activities_use_shared_gradebook_rules(self):
        ids, activities = self._calculated_quiz_fixture([Decimal(100)] * 10)
        sub = activities[0].template_subcomponent
        first = GradingTemplateDetail.objects.create(template_subcomponent=sub, code="A", name="A", weight_percentage=50)
        second = GradingTemplateDetail.objects.create(template_subcomponent=sub, code="B", name="B", weight_percentage=50)
        GradeActivity.objects.filter(pk=activities[0].pk).update(template_detail=first)
        GradeActivity.objects.filter(pk=activities[1].pk).update(template_detail=second)
        StudentActivityScore.objects.filter(activity=activities[1]).update(raw_score=48, computed_score=98)
        third = GradeActivity.objects.create(tenant=self.tenant, campus=self.campus, offering=self.offering,
            template_period=self.period, template_component=sub.template_component, template_subcomponent=sub,
            template_detail=second, title="Q3", total_score=50)
        for sid in ids:
            StudentActivityScore.objects.create(activity=third, student_id=sid, raw_score=46, computed_score=96)
        for mode, expected in (("WEIGHTED_DETAILS", "98.50"), ("AVERAGE_ACTIVITIES", "98.00")):
            with self.subTest(mode=mode):
                sub.detail_computation_mode = mode
                sub.save()
                self._persist_formula_totals(ids)
                response = self.search()
                quiz = self.detail(response).json()["quiz"]
                self.assertEqual(quiz["frequency"], [{"grade": expected, "count": 10, "percentage": "100.00"}])
                record = self._privacy_xlsx_row(self.client.get(self.url + "?" + response.context["export_query"]), "Quizzes")
                self.assertEqual(record["Exactly 98"], "10" if mode == "AVERAGE_ACTIVITIES" else "0")

    def test_quiz_average_across_major_components_keeps_decimal_precision(self):
        ids, _ = self._calculated_quiz_fixture([Decimal(100)] * 10)
        GradingTemplateComponent.objects.filter(template_period=self.period, is_exam_component=True).update(weight_percentage=20)
        other = GradingTemplateComponent.objects.create(template_period=self.period, code="EXTRA_QUIZZES", name="Quizzes", weight_percentage=20, score_input_mode="DIRECT_PERCENTAGE")
        activity = GradeActivity.objects.create(tenant=self.tenant, campus=self.campus, offering=self.offering,
            template_period=self.period, template_component=other, title="Other Quiz", total_score=100)
        for sid in ids:
            StudentActivityScore.objects.create(activity=activity, student_id=sid, raw_score=90, computed_score=90)
        self._persist_formula_totals(ids)
        response = self.search()
        quiz = self.detail(response).json()["quiz"]
        self.assertEqual(quiz["frequency"], [{"grade": "97.50", "count": 10, "percentage": "100.00"}])
        self.assertEqual(quiz["exact_100_count"], 0)
        record = self._privacy_xlsx_row(self.client.get(self.url + "?" + response.context["export_query"]), "Quizzes")
        self.assertTrue(all(record[f"Exactly {g}"] == "0" for g in range(100, 94, -1)))

    def test_quiz_missing_or_conflicting_classification_is_explicit_and_blank(self):
        ids, activities = self._calculated_quiz_fixture()
        sub = activities[0].template_subcomponent
        for code, name, reason in (("QUIZZES", "Written work", "disagree"), ("OTHER", "Quizzes", "disagree"), ("OTHER", "Written work", "no configured")):
            with self.subTest(code=code, name=name):
                sub.code, sub.name = code, name
                sub.save()
                response = self.search()
                quiz = self.detail(response).json()["quiz"]
                self.assertEqual(quiz["analytics_state"], "unavailable")
                self.assertIn(reason, quiz["calculation_note"])
                self.assertEqual(quiz["frequency"], [])
                self.assertTrue(all(self._privacy_xlsx_row(self.client.get(self.url + "?" + response.context["export_query"]), "Quizzes")[f"Exactly {g}"] == "" for g in range(100, 94, -1)))

    def test_quiz_historical_and_stale_lineage_are_not_finalized_snapshots(self):
        ids, activities = self._calculated_quiz_fixture()
        StudentPeriodGrade.objects.filter(offering=self.offering).update(is_finalized=True)
        current = self.detail(self.search(status="finalized")).json()["quiz"]
        self.assertEqual(current["analytics_state"], "disclosed")
        self.assertIn("not an immutable finalized snapshot", current["calculation_note"])
        StudentActivityScore.objects.filter(activity=activities[0]).update(computed_score=70)
        stale = self.detail(self.search(status="finalized")).json()["quiz"]
        self.assertEqual(stale["analytics_state"], "unavailable")
        self.assertIn("saved totals differ", stale["calculation_note"])
        self.offering.status = "ARCHIVED"
        self.offering.save()
        response = self.search(status="finalized")
        historical = self.detail(response).json()["quiz"]
        self.assertEqual(historical["analytics_state"], "unavailable")
        self.assertIn("historical", historical["calculation_note"])
        self.assertTrue(all(self._privacy_xlsx_row(self.client.get(self.url + "?" + response.context["export_query"]), "Quizzes")[f"Exactly {g}"] == "" for g in range(100, 94, -1)))

    def test_quiz_sparse_contributor_protects_aggregate_and_saved_ancestors(self):
        ids, activities = self._calculated_quiz_fixture([Decimal(100)] * 10)
        StudentActivityScore.objects.filter(activity=activities[1]).exclude(student_id=ids[0]).update(is_active=False)
        self._persist_formula_totals(ids)
        response = self.search()
        data = self.detail(response).json()
        for metric in (data, data["quiz"], data["measures"][1]):
            self._privacy_assert_payload(metric)
        self.assertEqual(data["measures"][0]["analytics_state"], "disclosed")
        records = self._assert_count_export_reconciliation(response)
        for record in records[1:]:
            self.assertTrue(all(record[f"Exactly {g}"] is None for g in range(100, 94, -1)))

    def test_quiz_complementary_statuses_nine_one_withhold_five_five_disclose(self):
        ids, _ = self._calculated_quiz_fixture([Decimal(100)] * 10)
        StudentPeriodGrade.objects.filter(offering=self.offering).update(is_finalized=True)
        StudentPeriodGrade.objects.filter(offering=self.offering, student_id=ids[-1]).update(is_finalized=False)
        for status in ("all", "finalized", "unfinalized"):
            response = self.search(status=status)
            self._privacy_assert_payload(self.detail(response).json()["quiz"])
            self._assert_count_export_reconciliation(response)
        StudentPeriodGrade.objects.filter(offering=self.offering, student_id__in=ids[5:]).update(is_finalized=False)
        for status, count in (("all", 10), ("finalized", 5), ("unfinalized", 5)):
            response = self.search(status=status)
            quiz = self.detail(response).json()["quiz"]
            self.assertEqual((quiz["analytics_state"], quiz["exact_100_count"]), ("disclosed", count))
            self._assert_count_export_reconciliation(response)

    def test_quiz_zero_through_six_distinct_student_boundary(self):
        from django.db import transaction
        ids, _ = self._calculated_quiz_fixture([Decimal(100)] * 10)
        for count in range(7):
            with self.subTest(count=count), transaction.atomic():
                Enrollment.objects.filter(course_offering=self.offering).exclude(student_id__in=ids[:count]).update(is_active=False)
                response = self.search()
                quiz = self.detail(response).json()["quiz"]
                state = "unavailable" if count == 0 else "privacy_suppressed" if count < 5 else "disclosed"
                self.assertEqual(quiz["analytics_state"], state)
                self.assertEqual(quiz["graded_count"], count)
                record = self._privacy_xlsx_row(self.client.get(self.url + "?" + response.context["export_query"]), "Quizzes")
                self.assertEqual(record["Exactly 100"], str(count) if count >= 5 else "")
                transaction.set_rollback(True)

    def test_quiz_overlap_cannot_release_small_population_difference(self):
        ids, activities = self._calculated_quiz_fixture([Decimal(100)] * 10)
        StudentActivityScore.objects.filter(activity=activities[0]).exclude(student_id__in=ids[:5]).update(is_active=False)
        StudentActivityScore.objects.filter(activity=activities[1]).exclude(student_id__in=ids[4:9]).update(is_active=False)
        self._persist_formula_totals(ids)
        response = self.search()
        quiz = self.detail(response).json()["quiz"]
        self._privacy_assert_payload(quiz)
        self._assert_count_export_reconciliation(response)

    def test_quiz_zero_weight_does_not_suppress_independent_saved_period(self):
        ids, activities = self._calculated_quiz_fixture([Decimal(100)] * 10)
        GradingTemplateComponent.objects.filter(pk=activities[0].template_component_id).update(weight_percentage=0)
        GradingTemplateComponent.objects.filter(template_period=self.period, is_exam_component=True).update(weight_percentage=100)
        StudentActivityScore.objects.filter(activity=activities[1]).exclude(student_id=ids[0]).update(is_active=False)
        self._persist_formula_totals(ids)
        response = self.search()
        data = self.detail(response).json()
        self._privacy_assert_payload(data["quiz"])
        self.assertEqual((data["analytics_state"], Decimal(data["average"])), ("disclosed", Decimal(96)))
        self.assertEqual(data["measures"][0]["analytics_state"], "disclosed")
        self._assert_count_export_reconciliation(response)

    def test_quiz_adapter_query_count_does_not_grow_per_student(self):
        from apps.admin_portal.grade_distribution_metrics import quiz_samples
        ids, activities = self._calculated_quiz_fixture([Decimal(100)] * 10)
        scores = list(StudentActivityScore.objects.filter(activity__in=activities).values("student_id", "activity_id", "computed_score"))
        counts = []
        for cohort in (set(ids[:3]), set(ids)):
            with CaptureQueriesContext(connection) as queries:
                values, sources, note = quiz_samples(self.period, activities, scores, cohort)
            self.assertEqual(len(values), len(cohort))
            self.assertEqual(set(values.values()), {Decimal(100)})
            self.assertEqual(sources, {a.id for a in activities})
            self.assertEqual(note, "")
            counts.append(len(queries))
        self.assertEqual(counts[0], counts[1])
        self.assertLessEqual(counts[1], 3)

    def test_export_has_exact_count_only_schema_and_integer_cells(self):
        self._calculated_quiz_fixture()
        response = self.search()
        sheet = self._workbook(self.client.get(self.url + "?" + response.context["export_query"])).active
        self.assertEqual([c.value for c in sheet[1]], ["Faculty", "Campus", "Course / Section", "Grading Period", "Computed Grade Component", "Exactly 100", "Exactly 99", "Exactly 98", "Exactly 97", "Exactly 96", "Exactly 95"])
        self.assertEqual(sheet.max_row, 4)
        self.assertEqual([sheet.cell(i, 5).value for i in (2, 3, 4)], ["Periodic Examination", "Class Standing", "Quizzes"])
        for cells in sheet.iter_rows(min_row=2, min_col=6):
            for cell in cells:
                self.assertIs(type(cell.value), int)
                self.assertEqual(cell.number_format, "0")
                self.assertIsNone(cell.comment)
        self.assertEqual(sheet.sheet_state, "visible")

    def test_export_uses_saved_buckets_without_normalizing_or_substituting_period_grade(self):
        StudentPeriodGrade.objects.filter(offering=self.offering).update(period_grade=100, exam_grade=150, class_standing_grade=99)
        self.client.force_login(self.admin_user)
        response = self.search()
        records = self._assert_count_export_reconciliation(response)
        self.assertTrue(all(records[0][f"Exactly {g}"] == 0 for g in range(100, 94, -1)))
        self.assertEqual(records[1]["Exactly 99"], 10)
        self.assertEqual(records[1]["Exactly 100"], 0)
        self.assertTrue(all(records[2][f"Exactly {g}"] is None for g in range(100, 94, -1)))


class GradeDistributionPrivacyMetricTests(SimpleTestCase):
    def assert_withheld(self, result):
        for key in ("average", "minimum", "maximum", "exact_100_count", "exact_100_percentage", "all_100"):
            self.assertIsNone(result[key], key)
        self.assertEqual(result["frequency"], [])
        self.assertTrue(all(item["count"] is None and item["percentage"] is None
                            for item in result["high_grade_breakdown"]))
        self.assertTrue(all(b["count"] is None and b["percentage"] is None for b in result["distribution"]))

    def test_exact_and_combined_high_grades_do_not_round_or_cap_values(self):
        values = list(map(Decimal, ["100", "100.01", "99.99", "99", "98", "97", "96", "95", "95.50", "90", "89.99", "0"]))
        result = metrics(values + [None], threshold=75, eligible_count=20)
        high = result["high_grade_breakdown"]
        self.assertEqual([item["count"] for item in high], [1, 1, 1, 1, 1, 1, 7, 10])
        self.assertTrue(all(item["total_graded_students"] == 12 for item in high))
        self.assertEqual(high[0]["percentage"], Decimal("8.33"))
        self.assertEqual(high[6]["percentage"], Decimal("58.33"))
        self.assertEqual(high[7]["percentage"], Decimal("83.33"))
        # No new information beyond the existing safely admitted frequency map.
        self.assertEqual(high[7]["count"], sum(item["count"] for item in result["frequency"] if item["grade"] >= 90))

    def test_high_grade_denominators_follow_each_measure_and_selected_status(self):
        statuses = {i: i < 5 for i in range(10)}
        samples = {i: Decimal("100") if i < 5 else Decimal("99.75") for i in statuses}
        for status, count, exact, combined in (("all", 10, 5, 5), ("finalized", 5, 5, 0), ("unfinalized", 5, 0, 5)):
            result = PrivacyFamily(statuses, status, 75).measure(samples, eligible_count=20)
            self.assertEqual(result["analytics_state"], "disclosed")
            high = result["high_grade_breakdown"]
            self.assertEqual((high[0]["count"], high[6]["count"], high[7]["count"]), (exact, combined, count))
            self.assertTrue(all(item["total_graded_students"] == count for item in high))
            self.assertEqual(high[7]["percentage"], Decimal("100.00"))

    def test_zero_to_six_distinct_saved_values_and_real_zero(self):
        for n in range(7):
            with self.subTest(n=n):
                privacy = PrivacyFamily({i: False for i in range(n)}, "all", Decimal("75"))
                result = privacy.measure({i: Decimal("0") for i in range(n)}, eligible_count=10)
                self.assertEqual(result["graded_count"], n)
                self.assertEqual(result["analytics_state"], "unavailable" if n == 0 else ("privacy_suppressed" if n < 5 else "disclosed"))
                if n < 5:
                    self.assert_withheld(result)
                else:
                    self.assertEqual(result["average"], 0)
                    self.assertEqual(result["minimum"], 0)

    def test_nine_finalized_one_unfinalized_suppresses_all_three_selections(self):
        statuses = {i: i < 9 for i in range(10)}
        for status in ("all", "finalized", "unfinalized"):
            result = PrivacyFamily(statuses, status, 75).measure({i: 97 for i in range(10)}, eligible_count=10)
            self.assertTrue(result["privacy_suppressed"])
            self.assert_withheld(result)
        self.assertEqual(PrivacyFamily(statuses, "finalized", 75).measure({i: 97 for i in range(10)}, eligible_count=10)["privacy_reason"], "complementary_population")

    def test_five_finalized_five_unfinalized_allow_all_and_both_subsets(self):
        statuses = {i: i < 5 for i in range(10)}
        for status, count in (("all", 10), ("finalized", 5), ("unfinalized", 5)):
            result = PrivacyFamily(statuses, status, 75).measure({i: 97 for i in range(10)}, eligible_count=10)
            self.assertEqual(result["analytics_state"], "disclosed")
            self.assertEqual(result["graded_count"], count)

    def test_owner_measure_example_period_thirty_exam_three_standing_twenty_five(self):
        family = PrivacyFamily({i: False for i in range(30)}, "all", 75)
        period = family.measure({i: 90 for i in range(30)}, eligible_count=30)
        exam = family.measure({i: 100 for i in range(3)}, eligible_count=30)
        standing = family.measure({i: 80 for i in range(25)}, eligible_count=30)
        self.assertEqual([r["analytics_state"] for r in (period, exam, standing)], ["disclosed", "privacy_suppressed", "disclosed"])
        self.assert_withheld(exam)

    def test_overlapping_populations_prevent_small_complements_and_intersections(self):
        family = PrivacyFamily({i: False for i in range(10)}, "all", 75)
        first = family.measure({i: 90 for i in range(10)}, eligible_count=10)
        second = family.measure({i: 80 for i in range(5)}, eligible_count=10)
        # All sets have at least five; comparing these two subsets would isolate one.
        third = family.measure({i: 70 for i in (0, 1, 2, 3, 5)}, eligible_count=10)
        self.assertEqual((first["analytics_state"], second["analytics_state"]), ("disclosed", "disclosed"))
        self.assertEqual(third["privacy_reason"], "complementary_population")
        self.assert_withheld(third)

    def test_unknown_period_status_is_a_third_complement_for_components(self):
        family = PrivacyFamily({i: True for i in range(9)}, "all", 75)
        result = family.measure({i: 90 for i in range(10)}, eligible_count=10)
        self.assertTrue(result["privacy_suppressed"])
        self.assert_withheld(result)

    def test_repeated_source_records_cannot_inflate_distinct_population(self):
        samples = {i % 4: 99 for i in range(40)}
        result = PrivacyFamily({i: False for i in range(4)}, "all", 75).measure(samples, eligible_count=10)
        self.assertEqual(result["graded_count"], 4)
        self.assert_withheld(result)

    def test_formula_plan_sparse_leaf_blocks_all_paths_not_just_memberships(self):
        family = PrivacyFamily({i: False for i in range(10)}, "all", 75)
        plan = FormulaPrivacyPlan(family)
        for key, samples in (("period", {i: 90 for i in range(10)}),
                             ("exam", {i: 79 for i in range(10)}),
                             ("common", {i: 75 for i in range(10)}), ("supplemental", {0: 43})):
            plan.add(key, samples)
        for a, b in (("period", "exam"), ("exam", "common"), ("exam", "supplemental")):
            plan.link(a, b)
        plan.finalize()
        for key in ("period", "exam", "common", "supplemental"):
            self.assert_withheld(plan.measure(key, eligible_count=10))

    def test_formula_plan_unrelated_family_and_unknown_dependencies(self):
        for unresolved in (False, True):
            family = PrivacyFamily({i: False for i in range(10)}, "all", 75)
            plan = FormulaPrivacyPlan(family)
            plan.add("safe_period", {i: 90 for i in range(10)})
            plan.add("exam", {i: 79 for i in range(10)})
            plan.add("supplemental", {0: 43})
            plan.link("exam", "supplemental")
            plan.finalize(unresolved=unresolved)
            period = plan.measure("safe_period", eligible_count=10)
            self.assertEqual(period["analytics_state"], "privacy_suppressed" if unresolved else "disclosed")
            self.assert_withheld(plan.measure("exam", eligible_count=10))

    def test_formula_plan_never_calculates_metrics_for_blocked_family(self):
        plan = FormulaPrivacyPlan(PrivacyFamily({i: False for i in range(10)}, "all", 75))
        plan.add("period", {i: 90 for i in range(10)})
        plan.add("source", {0: 43})
        plan.link("period", "source")
        with patch("apps.admin_portal.grade_distribution_metrics.metrics", side_effect=AssertionError("hidden statistics calculated")):
            plan.finalize()
            self.assert_withheld(plan.measure("period", eligible_count=10))

    def test_formula_plan_retains_safe_siblings_after_cutting_all_ancestor_paths(self):
        plan = FormulaPrivacyPlan(PrivacyFamily({i: False for i in range(10)}, "all", 75))
        for key, samples in (("period", {i: 90 for i in range(10)}), ("exam", {i: 79 for i in range(10)}),
                             ("common", {i: 75 for i in range(10)}), ("standing", {i: 80 for i in range(10)}),
                             ("supplemental", {0: 43})):
            plan.add(key, samples)
        for parent, child in (("period", "exam"), ("period", "standing"), ("exam", "common"), ("exam", "supplemental")):
            plan.depends_on(parent, child)
        plan.finalize()
        for key in ("period", "exam", "supplemental"):
            self.assert_withheld(plan.measure(key, eligible_count=10))
        for key in ("common", "standing"):
            self.assertEqual(plan.measure(key, eligible_count=10)["analytics_state"], "disclosed")
