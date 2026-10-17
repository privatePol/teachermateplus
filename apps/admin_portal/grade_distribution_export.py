"""Bounded count-only export of authorized, internal Gradebook aggregates."""
from collections import defaultdict
from io import BytesIO

from django.core.exceptions import ValidationError
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

MAX_WORKBOOK_ROWS = 50000
HEADERS = ["Faculty", "Campus", "Course / Section", "Grading Period", "Computed Grade Component"]


def validate_workbook_size(count):
    if count > MAX_WORKBOOK_ROWS:
        raise ValidationError(f"Excel export limit is {MAX_WORKBOOK_ROWS:,} aggregate rows per sheet. Narrow your criteria; no records were truncated.")


def export_components(row):
    # Saved totals are stable machine sources, never matched by user labels.
    by_source = {m["dataset_id"].split("/")[-2]: m for m in row["measures"] if m["activity_id"] is None}
    return [("Period Grades", row),
            ("Examination Grades", by_source["exam_grade"]),
            ("Class Standing Grades", by_source["class_standing_grade"]),
            ("Quizzes", row["quiz"])]


def class_labels(rows):
    """Qualify only colliding authorized classes; never group their counts.

    Department/program describe the offering, not its students or the faculty's
    supervision department. Final visible identities must distinguish physical
    datasets, including when authored labels collide with a qualifier.
    """
    def normalized(*values):
        return tuple(" ".join(str(value).split()).casefold() for value in values)

    labels = [[f'{r["course_code"]} - {r["course_title"]} / {r["section"]}', r["period"]] for r in rows]
    classes, periods = defaultdict(list), defaultdict(list)
    for index, row in enumerate(rows):
        classes[normalized(row["faculty_name"], row["campus"], *labels[index])].append(index)
        periods[(row["offering_id"], normalized(row["period"]))].append(index)

    for indexes in classes.values():
        if len({rows[i]["offering_id"] for i in indexes}) < 2:
            continue
        if len({normalized(rows[i]["department"]) for i in indexes}) > 1:
            for i in indexes:
                labels[i][0] += f' [Department: {rows[i]["department"]}]'
        remaining = defaultdict(list)
        for i in indexes:
            remaining[normalized(labels[i][0])].append(i)
        for candidates in remaining.values():
            if (len({rows[i]["offering_id"] for i in candidates}) > 1
                    and len({normalized(rows[i]["program"]) for i in candidates}) > 1):
                for i in candidates:
                    labels[i][0] += f' [Program: {rows[i]["program"] or "Not specified"}]'

    for indexes in periods.values():
        if len({rows[i]["physical_period_id"] for i in indexes}) > 1:
            for i in indexes:
                labels[i][1] += f' ({rows[i]["period_context"]})'

    visible = defaultdict(set)
    for row, label in zip(rows, labels):
        visible[normalized(row["faculty_name"], row["campus"], *label)].add(
            (row["offering_id"], row["physical_period_id"]))
    if any(len(datasets) > 1 for datasets in visible.values()):
        raise ValidationError(
            "Excel class labels remain ambiguous after department/program qualification. "
            "Correct the class identifying data before exporting; no records were exported.")
    return labels


def grade_distribution_workbook(rows, *, range_heading=None):
    validate_workbook_size(4 * len(rows))
    labels = class_labels(rows)
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Grade Distribution"
    label = range_heading or (rows[0]["range_label"] if rows else "Grade Range")
    sheet.append(HEADERS + [label, "Total Graded Students", "Percentage", "Analytics Status", "Notice"])
    for row, (course, period) in zip(rows, labels):
        common = [row["faculty_name"], row["campus"], course, period]
        for label, measure in export_components(row):
            disclosed = measure.get("range_state") == "disclosed"
            values = [measure["range_count"], measure["range_denominator"], measure["range_percentage"]] if disclosed else [None] * 3
            notice = " ".join(filter(None, (measure.get("privacy_message"), measure.get("source_notice"), measure.get("history_notice"), measure.get("calculation_note"), measure.get("unit"))))
            sheet.append(common + [label] + values + ["Available" if disclosed else "Unavailable", notice])
    sheet.freeze_panes = "F2"
    sheet.auto_filter.ref = sheet.dimensions
    sheet.sheet_view.showGridLines = False
    sheet.print_title_rows = "1:1"
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    sheet.page_setup.orientation = "landscape"
    sheet.page_setup.paperSize = sheet.PAPERSIZE_A4
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 0
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="20613C")
        cell.alignment = Alignment(wrap_text=True, vertical="center")
        sheet.column_dimensions[cell.column_letter].width = 40 if cell.column <= 5 else 14
    sheet.row_dimensions[1].height = 32
    for cells in sheet.iter_rows(min_row=2):
        for cell in cells:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            if isinstance(cell.value, str):
                # User-authored labels, including = + - @, are literal text.
                cell.data_type = "s"
            elif cell.value is not None:
                cell.number_format = "0.00" if cell.column == 8 else "0"
            if cell.row % 2 == 0:
                cell.fill = PatternFill("solid", fgColor="F0F7F2")
    output = BytesIO()
    workbook.save(output)
    return output.getvalue()
