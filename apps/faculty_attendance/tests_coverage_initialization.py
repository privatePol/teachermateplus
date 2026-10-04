"""Focused legacy activation, frozen attribution recovery and paper/export cases."""
from copy import deepcopy
from datetime import date
from io import BytesIO

from django.core.exceptions import PermissionDenied, ValidationError
from django.test import Client, TestCase
from django.urls import reverse
from openpyxl import load_workbook

from apps.academics.models import FacultyAssignment
from apps.rbac.models import Permission, UserPermission
from . import tests as existing
from .coverage_initialization import CoverageAdoptionService, CoverageInitializationService
from .cutoffs import review_cutoff, publish_cutoff
from .models import (AttendanceResult, FacultyCoverage, MeetingCoverageAdoption, CoverageReconciliation,
                     AttendanceClosureDecision, AttendanceCutoffPublication, AttendanceCutoffPublicationEntry)
from .monthly_checklists import MonthlyArrangementService
from .observations import AttendanceResultService, CheckingRoundService, ObservationService, require_confirmable_meeting_faculty
from .permissions import MANAGE_COVERAGE_PERMISSION, RECONCILE_PERMISSION, PRINT_PERMISSION
from .selectors import unresolved_meetings, faculty_meeting_history
from .services import MeetingService, SubstitutionService, ReconciliationService, _meeting_snapshot


class CoverageInitializationTests(TestCase):
    permission_codes = existing.FacultyAttendanceFoundationTests.permission_codes
    setUpTestData = classmethod(existing.FacultyAttendanceFoundationTests.setUpTestData.__func__)
    setUp = existing.FacultyAttendanceFoundationTests.setUp
    aware = existing.FacultyAttendanceFoundationTests.aware
    schedule = existing.FacultyAttendanceFoundationTests.schedule
    coverage = existing.FacultyAttendanceFoundationTests.coverage
    meeting = existing.FacultyAttendanceFoundationTests.meeting

    def assignment(self, offering=None, faculty=None, **kwargs):
        return FacultyAssignment.objects.create(offering=offering or self.offering,
            faculty_user=faculty or self.faculty, is_primary=True, response_status='ACCEPTED',
            accepted_at=self.aware(2026, 1, 4), **kwargs)

    def scope(self, **kwargs):
        return dict(actor=self.actor, tenant_id=self.tenant.pk, campus_id=self.campus.pk,
            academic_year=self.academic_year, term=self.term,
            effective_from=kwargs.pop('effective_from', self.aware(2026, 1, 1)), **kwargs)

    def apply(self, recover=False, **kwargs):
        scope = self.scope(**kwargs)
        preview = CoverageInitializationService.preview(**scope)
        return CoverageInitializationService.apply(**scope, fingerprint=preview['fingerprint'],
            recover=recover, confirmed=True, reason='Verified synthetic teaching boundary')

    def unresolved(self, combined=False):
        slot = self.schedule().slots.get()
        return MeetingService.generate(actor=self.actor, schedule_slot=slot, meeting_date=date(2026, 1, 5),
            offerings=[self.combined_offering] if combined else [])

    def confirm(self, meeting, round_=None):
        round_ = round_ or CheckingRoundService.create(actor=self.actor, meetings=[meeting], checking_date=meeting.meeting_date)
        AttendanceResultService.confirm_present(actor=self.actor, checking_round=round_, manifest_revision=1,
            reviewed_rows=[{'meeting_id': meeting.pk, 'result_revision': 0}])
        return AttendanceResult.objects.get(meeting=meeting)

    def test_null_assignment_scope_uses_owning_offering_and_explicit_boundary_idempotently(self):
        assignment = self.assignment()
        self.faculty.default_tenant = self.faculty.default_campus = self.faculty.default_department = None
        self.faculty.save()
        result = self.apply(effective_from=self.aware(2026, 1, 5, 8))
        self.assertEqual(result['created'], 1)
        coverage = FacultyCoverage.objects.get(offering=self.offering)
        self.assertEqual(coverage.source_assignment, assignment)
        self.assertEqual(coverage.effective_from, self.aware(2026, 1, 5, 8))
        self.assertEqual(coverage.effective_until, self.aware(2026, 6, 1))
        self.assertIsNone(existing.CoverageService.effective_for(offering_id=self.offering.pk, at=self.aware(2026, 1, 5, 7)))
        self.assertEqual(self.apply(effective_from=self.aware(2026, 1, 5, 8))['created'], 0)
        self.assertEqual(FacultyCoverage.objects.count(), 1)

    def test_multiple_unaccepted_and_explicit_scope_conflicts_are_review_only(self):
        assignment = self.assignment()
        assignment.campus = existing.Campus.objects.create(tenant=self.tenant, code='OTHER', name='Other')
        assignment.save()
        self.assertEqual(self.apply()['created'], 0)
        assignment.campus = None
        assignment.response_status = 'PENDING'
        assignment.save()
        self.assertEqual(self.apply()['created'], 0)
        assignment.response_status = 'ACCEPTED'
        assignment.save()
        self.assignment(faculty=self.replacement, is_active=True)
        self.assertEqual(self.apply()['created'], 0)

    def test_overlap_is_preserved_not_replaced_or_extended(self):
        self.assignment()
        coverage = self.coverage(faculty=self.replacement)
        self.assertEqual(self.apply()['created'], 0)
        coverage.refresh_from_db()
        self.assertEqual(coverage.faculty_user, self.replacement)
        self.assertIsNone(coverage.effective_until)

    def test_preview_staleness_confirmation_and_boundary_validation(self):
        assignment = self.assignment()
        scope = self.scope()
        plan = CoverageInitializationService.preview(**scope)
        with self.assertRaises(ValidationError):
            CoverageInitializationService.apply(**scope, fingerprint=plan['fingerprint'], recover=False, confirmed=False)
        assignment.is_active = False
        assignment.save()
        with self.assertRaises(ValidationError):
            CoverageInitializationService.apply(**scope, fingerprint=plan['fingerprint'], recover=False, confirmed=True)
        with self.assertRaises(ValidationError):
            CoverageInitializationService.preview(**self.scope(effective_from=self.aware(2025, 12, 31)))
        self.assertFalse(FacultyCoverage.objects.exists())

    def test_initial_assignment_queue_is_resolved_but_replacement_is_not(self):
        assignment = self.assignment()
        queue = CoverageReconciliation.objects.create(tenant=self.tenant, campus=self.campus,
            department=self.department, offering=self.offering, source_assignment=assignment,
            proposed_faculty=self.faculty, created_by=self.actor, event_type='ASSIGNMENT_ACCEPTED',
            source_reference='synthetic-initial')
        self.assertEqual(self.apply()['created'], 1)
        queue.refresh_from_db()
        self.assertEqual(queue.status, 'RESOLVED')
        self.assertEqual(queue.effective_at, self.aware(2026, 1, 1))
        queue.pk = None
        queue.source_reference = 'synthetic-replacement'
        queue.event_type, queue.prior_faculty = 'PERMANENT_REPLACEMENT', self.replacement
        queue.status, queue.resolved_by, queue.resolved_at, queue.resolution_reason = 'PENDING', None, None, ''
        queue.save()
        self.assertEqual(CoverageInitializationService.preview(**self.scope())['rows'][0]['status'], 'REVIEW')

    def test_frozen_recovery_confirm_publication_and_monitoring_share_adoption(self):
        self.assignment()
        meeting = self.unresolved()
        round_ = CheckingRoundService.create(actor=self.actor, meetings=[meeting], checking_date=meeting.meeting_date)
        frozen = deepcopy(round_.manifest_rows.get().meeting_snapshot)
        original = deepcopy(meeting.schedule_snapshot)
        original_meeting = deepcopy(_meeting_snapshot(meeting))
        outcome = self.apply(recover=True)
        self.assertEqual((outcome['created'], outcome['adopted'], outcome['recovery']), (1, 1, []))
        self.assertTrue(existing.AuditLog.objects.filter(action='FACULTY_ATTENDANCE_COVERAGE_ADOPTED').exists())
        meeting.refresh_from_db()
        self.assertTrue(meeting.unresolved_coverage)
        self.assertIsNone(meeting.faculty_user_id)
        self.assertEqual(meeting.schedule_snapshot, original)
        self.assertEqual(_meeting_snapshot(meeting), original_meeting)
        self.assertEqual(round_.manifest_rows.get().meeting_snapshot, frozen)
        self.assertEqual(require_confirmable_meeting_faculty(meeting), self.faculty)
        self.assertFalse(unresolved_meetings(tenant_id=self.tenant.pk).exists())
        self.assertIn(meeting, faculty_meeting_history(tenant_id=self.tenant.pk, faculty_user_id=self.faculty.pk))
        result = self.confirm(meeting, round_)
        self.assertEqual((result.faculty_user, result.revision, result.status), (self.faculty, 1, 'PRESENT'))
        # A second active Monday class would intentionally make this campus incomplete.
        self.combined_offering.schedule_text = 'T 08:00-09:00'
        self.combined_offering.save(update_fields=['schedule_text'])
        review = review_cutoff(actor=self.actor, tenant_id=self.tenant.pk, campus_id=self.campus.pk,
            academic_year=self.academic_year, term=self.term, start_date=meeting.meeting_date, end_date=meeting.meeting_date)
        self.assertFalse(review.blockers, review.blockers)
        publication = publish_cutoff(actor=self.actor, tenant_id=self.tenant.pk, campus_id=self.campus.pk,
            academic_year=self.academic_year, term=self.term, start_date=meeting.meeting_date, end_date=meeting.meeting_date,
            expected_fingerprint=review.fingerprint, submission_key='adoption-publication', publication_reason='')
        entry = publication.entries.get()
        self.assertEqual(entry.faculty_user, self.faculty)
        self.assertIn('coverage_adoption', entry.meeting_snapshot)
        from .monitoring import term_summary
        report = term_summary(actor=self.actor, tenant_id=self.tenant.pk, campus_id=self.campus.pk,
            academic_year=self.academic_year, term=self.term, as_of=meeting.meeting_date)
        self.assertEqual(report['rows'][0]['actual'], existing.Decimal('1.00'))

    def test_adopted_card_and_exception_save_use_new_attribution_not_original_flag(self):
        self.assignment()
        meeting = self.unresolved()
        round_ = CheckingRoundService.create(actor=self.actor, meetings=[meeting], checking_date=meeting.meeting_date)
        self.apply(recover=True)
        self.client.force_login(self.actor)
        response = self.client.get(reverse('faculty_attendance:round', args=[round_.public_id]))
        self.assertContains(response, 'Checker-approved dated coverage')
        self.assertContains(response, 'data-row-save')
        self.assertNotContains(response, 'This row cannot be saved')
        observation = ObservationService.record(actor=self.actor, checking_round=round_, meeting_id=meeting.pk,
            manifest_revision=1, submission_key='adoption-exception', findings=[{'finding_type': 'LATE', 'segment_key': 'arrival', 'minutes': 2}])
        result = AttendanceResultService.select_observation(actor=self.actor, observation=observation, expected_revision=0, reason='')
        self.assertEqual((result.faculty_user, result.late_minutes), (self.faculty, 2))

    def test_combined_adoption_requires_same_coverage_and_counts_one_meeting(self):
        self.assignment()
        self.assignment(offering=self.combined_offering)
        meeting = self.unresolved(combined=True)
        self.assertEqual(self.apply(recover=True)['adopted'], 1)
        self.assertEqual(MeetingCoverageAdoption.objects.count(), 1)
        self.assertEqual(meeting.offering_links.count(), 2)
        self.assertEqual(self.confirm(meeting).faculty_user, self.faculty)

    def test_conflicting_combined_substitution_and_recorded_findings_not_reassigned(self):
        self.assignment()
        self.assignment(offering=self.combined_offering, faculty=self.replacement)
        meeting = self.unresolved(combined=True)
        outcome = self.apply(recover=True)
        self.assertEqual(outcome['adopted'], 0)
        self.assertIn('conflicting', outcome['recovery'][0]['message'])
        with self.assertRaises(ValidationError):
            require_confirmable_meeting_faculty(meeting)

    def test_substitution_and_saved_history_are_preserved_and_adoption_is_immutable(self):
        self.assignment()
        meeting = self.unresolved()
        SubstitutionService.assign(actor=self.actor, meeting=meeting, substitute_faculty=self.replacement, reason='')
        self.assertEqual(self.apply(recover=True)['adopted'], 0)
        # New permanent coverage does not silently settle an explicit substitute's history review.
        for item in meeting.reconciliations.filter(status='PENDING'):
            ReconciliationService.resolve(actor=self.actor, reconciliation=item, decision='KEEP_SNAPSHOT', reason='')
        result = self.confirm(meeting)
        before = list(result.history.values())
        with self.assertRaises(ValidationError):
            CoverageAdoptionService.adopt(actor=self.actor, meeting=meeting)
        result.refresh_from_db()
        self.assertEqual(result.faculty_user, self.replacement)
        self.assertEqual(list(result.history.values()), before)

    def test_direct_deny_rejects_initialization_get_post_and_service(self):
        self.assignment()
        self.client.force_login(self.actor)
        UserPermission.objects.create(user=self.actor, permission=Permission.objects.get(code=MANAGE_COVERAGE_PERMISSION),
            grant_type='DENY', tenant=None, campus=None)
        for method in (self.client.get, self.client.post):
            self.assertEqual(method(reverse('faculty_attendance:coverage_initialize')).status_code, 403)
        with self.assertRaises(PermissionDenied):
            self.apply()
        self.assertFalse(FacultyCoverage.objects.exists())

    def test_scoped_preview_apply_requires_signed_unchanged_confirmation(self):
        self.assignment()
        self.client.force_login(self.actor)
        url = reverse('faculty_attendance:coverage_initialize')
        values = {'academic_year': self.academic_year.pk, 'term': self.term.pk, 'effective_from': '2026-01-05T08:00'}
        preview = self.client.post(url, {**values, 'action': 'preview'})
        self.assertEqual(preview.status_code, 200)
        self.assertFalse(FacultyCoverage.objects.exists())
        token = preview.context['preview_token']
        invalid = self.client.post(url, {**values, 'action': 'apply', 'preview_token': token, 'confirmed': 'on', 'effective_from': '2026-01-06T08:00'})
        self.assertContains(invalid, 'Scope or effective boundary changed')
        self.assertFalse(FacultyCoverage.objects.exists())
        valid = self.client.post(url, {**values, 'action': 'apply', 'preview_token': token, 'confirmed': 'on'})
        self.assertContains(valid, 'Created 1 coverage interval')
        self.assertEqual(FacultyCoverage.objects.count(), 1)

    def test_new_confirmed_initial_activation_does_not_modify_snapshot_without_adoption(self):
        self.assignment()
        meeting = self.unresolved()
        self.apply()
        self.assertFalse(MeetingCoverageAdoption.objects.exists())
        with self.assertRaises(ValidationError):
            require_confirmable_meeting_faculty(meeting)
        self.client.force_login(self.actor)
        url = reverse('faculty_attendance:reconciliation')
        response = self.client.post(url, {'kind': 'adoption', 'meeting_id': meeting.pk, 'confirmed': 'on',
                                         'reason': 'Synthetic verified legacy attribution'})
        self.assertEqual(response.status_code, 302)
        self.assertEqual(MeetingCoverageAdoption.objects.count(), 1)
        adoption = MeetingCoverageAdoption.objects.get()
        with self.assertRaises(ValidationError):
            adoption.save()
        with self.assertRaises(ValidationError):
            adoption.delete()

    def test_saved_initial_finding_and_other_history_conflict_cannot_be_adopted(self):
        self.assignment()
        meeting = self.unresolved()
        self.apply()
        # Represents a legacy recorded row. The recovery must not silently take ownership.
        saved = AttendanceResult.objects.create(meeting=meeting, faculty_user=self.replacement,
            revision=1, status='PRESENT')
        with self.assertRaisesMessage(ValidationError, 'Recorded or published'):
            CoverageAdoptionService.adopt(actor=self.actor, meeting=meeting)
        saved.refresh_from_db()
        self.assertEqual((saved.faculty_user, saved.revision), (self.replacement, 1))
        self.assertFalse(MeetingCoverageAdoption.objects.exists())

    def test_recovery_direct_deny_and_cross_campus_coverage_validation(self):
        assignment = self.assignment()
        meeting = self.unresolved()
        self.apply()
        UserPermission.objects.create(user=self.actor, permission=Permission.objects.get(code=RECONCILE_PERMISSION),
            grant_type='DENY', tenant=self.tenant, campus=self.campus)
        with self.assertRaises(PermissionDenied):
            CoverageAdoptionService.adopt(actor=self.actor, meeting=meeting)
        self.assertFalse(MeetingCoverageAdoption.objects.exists())
        assignment.campus = existing.Campus.objects.create(tenant=self.tenant, code='OTHER', name='Other')
        assignment.save()
        coverage = FacultyCoverage.objects.get()
        with self.assertRaisesMessage(ValidationError, 'Source assignment scope conflicts'):
            coverage.full_clean()

    def test_initialization_token_does_not_replace_csrf(self):
        client = Client(enforce_csrf_checks=True)
        client.force_login(self.actor)
        response = client.post(reverse('faculty_attendance:coverage_initialize'), {
            'action': 'apply', 'confirmed': 'on', 'preview_token': 'invented'})
        self.assertEqual(response.status_code, 403)
        self.assertFalse(FacultyCoverage.objects.exists())


class LegacyMeetingRecoveryUITests(TestCase):
    """Exceptional recovery UI; every record below is disposable synthetic data."""
    permission_codes = CoverageInitializationTests.permission_codes
    setUpTestData = classmethod(CoverageInitializationTests.setUpTestData.__func__)
    setUp = CoverageInitializationTests.setUp
    aware = CoverageInitializationTests.aware
    schedule = CoverageInitializationTests.schedule
    coverage = CoverageInitializationTests.coverage
    unresolved = CoverageInitializationTests.unresolved

    def ready(self, combined=False):
        meeting = self.unresolved(combined=combined)
        self.coverage()
        if combined:
            self.coverage(offering=self.combined_offering)
        for item in meeting.reconciliations.filter(status='PENDING'):
            ReconciliationService.resolve(actor=self.actor, reconciliation=item,
                decision='KEEP_SNAPSHOT', reason='Synthetic legacy history preserved')
        self.client.force_login(self.actor)
        return meeting

    def post(self, meeting, **changes):
        values = dict(kind='adoption', meeting_id=meeting.pk, confirmed='on',
                      reason='Synthetic verified legacy recovery')
        values.update(changes)
        return self.client.post(reverse('faculty_attendance:reconciliation'), values)

    def test_collapsed_panel_shows_verified_dated_candidate_and_all_linked_sections_read_only(self):
        meeting = self.ready(combined=True)
        before = deepcopy(_meeting_snapshot(meeting))
        audits = existing.AuditLog.objects.count()
        response = self.client.get(reverse('faculty_attendance:reconciliation'))
        self.assertContains(response, 'id="legacy-meeting-recovery">')
        self.assertContains(response, f'Meeting {meeting.pk}')
        self.assertContains(response, 'MATH101 (S1)')
        self.assertContains(response, 'MATH101 (S2)')
        self.assertContains(response, 'January 5, 2026')
        self.assertContains(response, '8:00 AM–9:00 AM')
        self.assertContains(response, 'Asia/Manila')
        self.assertEqual(response.context['adoption_items'][0].recovery_candidate, self.faculty)
        self.assertContains(response, 'Verified candidate:')
        self.assertContains(response, f'data-legacy-adoption="{meeting.pk}"')
        self.assertContains(response, 'name="confirmed"')
        self.assertContains(response, 'name="reason"')
        self.assertContains(response, 'name="csrfmiddlewaretoken"')
        self.assertNotContains(response, reverse('faculty_attendance:coverage_initialize'))
        self.assertFalse(MeetingCoverageAdoption.objects.exists())
        self.assertEqual(existing.AuditLog.objects.count(), audits)
        meeting.refresh_from_db()
        self.assertEqual(_meeting_snapshot(meeting), before)

    def test_missing_or_conflicting_evidence_is_not_presented_as_verified(self):
        meeting = self.unresolved(combined=True)
        self.client.force_login(self.actor)
        response = self.client.get(reverse('faculty_attendance:reconciliation'))
        self.assertContains(response, 'Candidate not verified')
        self.assertNotContains(response, 'Verified candidate:')
        self.assertNotContains(response, 'data-legacy-adoption=')
        self.assertContains(self.post(meeting), 'one verified coverage interval at class start')
        self.coverage()
        self.coverage(offering=self.combined_offering, faculty=self.replacement)
        response = self.client.get(reverse('faculty_attendance:reconciliation'))
        self.assertContains(response, 'conflicting dated faculty coverage')
        self.assertNotContains(response, 'data-legacy-adoption=')
        self.assertContains(self.post(meeting), 'conflicting dated faculty coverage')
        self.assertFalse(MeetingCoverageAdoption.objects.exists())

    def test_manage_coverage_and_reconciliation_direct_denies_hide_or_reject_recovery(self):
        meeting = self.ready()
        deny = UserPermission.objects.create(user=self.actor,
            permission=Permission.objects.get(code=MANAGE_COVERAGE_PERMISSION),
            grant_type='DENY', tenant=None, campus=None)
        response = self.client.get(reverse('faculty_attendance:reconciliation'))
        self.assertEqual(response.status_code, 200)
        self.assertNotContains(response, 'id="legacy-meeting-recovery"')
        self.assertEqual(self.post(meeting).status_code, 403)
        with self.assertRaises(PermissionDenied):
            CoverageAdoptionService.preview(actor=self.actor, meeting=meeting)
        deny.delete()
        UserPermission.objects.create(user=self.actor, permission=Permission.objects.get(code=RECONCILE_PERMISSION),
            grant_type='DENY', tenant=self.tenant, campus=self.campus)
        self.assertEqual(self.client.get(reverse('faculty_attendance:reconciliation')).status_code, 403)
        self.assertEqual(self.post(meeting).status_code, 403)
        self.assertFalse(MeetingCoverageAdoption.objects.exists())

    def test_linked_offering_department_scope_and_cross_campus_meeting_cannot_be_recovered(self):
        meeting = self.ready(combined=True)
        self.combined_offering.department = self.other_department
        self.combined_offering.save(update_fields=['department'])
        response = self.client.get(reverse('faculty_attendance:reconciliation'))
        self.assertNotContains(response, f'data-legacy-adoption="{meeting.pk}"')
        self.assertEqual(self.post(meeting).status_code, 403)
        other_campus = existing.Campus.objects.create(tenant=self.tenant, code='OTHER', name='Other')
        # Disposable legacy anomaly must not widen the scoped view's lookup.
        type(meeting).objects.filter(pk=meeting.pk).update(campus=other_campus)
        self.assertEqual(self.post(meeting).status_code, 404)
        self.assertFalse(MeetingCoverageAdoption.objects.exists())

    def test_csrf_required_even_for_confirmed_authorized_recovery(self):
        meeting = self.ready()
        client = Client(enforce_csrf_checks=True)
        client.force_login(self.actor)
        url = reverse('faculty_attendance:reconciliation')
        values = dict(kind='adoption', meeting_id=meeting.pk, confirmed='on', reason='Synthetic CSRF recovery')
        self.assertEqual(client.post(url, values).status_code, 403)
        self.assertFalse(MeetingCoverageAdoption.objects.exists())
        self.assertEqual(client.get(url).status_code, 200)
        values['csrfmiddlewaretoken'] = client.cookies['csrftoken'].value
        self.assertEqual(client.post(url, values).status_code, 302)
        self.assertEqual(MeetingCoverageAdoption.objects.get().reason, 'Synthetic CSRF recovery')
        self.assertTrue(existing.AuditLog.objects.filter(action='FACULTY_ATTENDANCE_COVERAGE_ADOPTED').exists())

    def test_reason_and_explicit_confirmation_required_with_retained_input(self):
        meeting = self.ready()
        for values, field in (({'reason': ''}, 'reason'), ({'reason': '   '}, 'reason'),
                              ({'confirmed': ''}, 'confirmed'), ({'confirmed': 'true'}, 'confirmed')):
            with self.subTest(values=values):
                response = self.post(meeting, **values)
                self.assertEqual(response.status_code, 200)
                self.assertIn(field, response.context['adoption_form_errors'])
                self.assertContains(response, 'id="legacy-meeting-recovery" open')
                self.assertContains(response, 'Recovery was not saved.')
                if field == 'confirmed':
                    self.assertContains(response, 'Synthetic verified legacy recovery')
                self.assertFalse(MeetingCoverageAdoption.objects.exists())

    def test_ge109_http_adoption_retains_boundaries_frozen_evidence_and_keep_snapshot(self):
        self.academic_year.start_date, self.academic_year.end_date = date(2026, 6, 1), date(2027, 5, 31)
        self.academic_year.save(update_fields=['start_date', 'end_date'])
        self.term.start_date, self.term.end_date = date(2026, 9, 1), date(2026, 12, 31)
        self.term.save(update_fields=['start_date', 'end_date'])
        self.offering.schedule_text = 'M 07:30-09:00'
        self.offering.save(update_fields=['schedule_text'])
        slot = self.schedule(effective_from=date(2026, 10, 2)).slots.get()
        meeting = MeetingService.generate(actor=self.actor, schedule_slot=slot, meeting_date=date(2026, 10, 5), offerings=[])
        round_ = CheckingRoundService.create(actor=self.actor, meetings=[meeting], checking_date=meeting.meeting_date)
        boundary = self.aware(2026, 10, 5, 8)
        first = self.coverage(effective_from=self.aware(2026, 10, 2, 8), effective_until=boundary)
        second = self.coverage(effective_from=boundary)
        for item in meeting.reconciliations.filter(status='PENDING'):
            ReconciliationService.resolve(actor=self.actor, reconciliation=item,
                decision='KEEP_SNAPSHOT', reason='Synthetic prior review retained')
        before = deepcopy(_meeting_snapshot(meeting))
        manifest = deepcopy(round_.manifest_rows.get().meeting_snapshot)
        reviews = list(meeting.reconciliations.values())
        intervals = list(FacultyCoverage.objects.values())
        self.client.force_login(self.actor)
        preview = self.client.get(reverse('faculty_attendance:reconciliation'))
        self.assertContains(preview, 'October 5, 2026')
        self.assertContains(preview, '7:30 AM–9:00 AM')
        self.assertEqual(self.post(meeting).status_code, 302)
        adoption = MeetingCoverageAdoption.objects.get()
        self.assertEqual((adoption.coverage, adoption.faculty_user, adoption.adopted_by), (first, self.faculty, self.actor))
        self.assertEqual(adoption.decision_snapshot['original_meeting'], before)
        meeting.refresh_from_db()
        self.assertEqual(require_confirmable_meeting_faculty(meeting), self.faculty)
        self.assertEqual(_meeting_snapshot(meeting), before)
        self.assertEqual(round_.manifest_rows.get().meeting_snapshot, manifest)
        self.assertEqual(list(meeting.reconciliations.values()), reviews)
        self.assertEqual(list(FacultyCoverage.objects.values()), intervals)
        self.assertEqual((first.effective_until, second.effective_from), (boundary, boundary))
        self.assertFalse(AttendanceResult.objects.exists())
        self.assertFalse(AttendanceCutoffPublication.objects.exists())
        audit = existing.AuditLog.objects.get(action='FACULTY_ATTENDANCE_COVERAGE_ADOPTED')
        self.assertEqual(audit.actor_user_id, self.actor.pk)
        # Existing service retry stays idempotent; the stale UI endpoint does not resubmit it.
        self.assertEqual(CoverageAdoptionService.adopt(actor=self.actor, meeting=meeting).pk, adoption.pk)
        self.assertEqual(self.post(meeting).status_code, 404)
        self.assertEqual(MeetingCoverageAdoption.objects.count(), 1)

    def test_saved_or_substituted_meetings_are_not_recovery_targets(self):
        meeting = self.ready()
        saved = AttendanceResult.objects.create(meeting=meeting, faculty_user=self.replacement,
            revision=1, status='PRESENT')
        history = list(saved.history.values())
        response = self.client.get(reverse('faculty_attendance:reconciliation'))
        self.assertNotContains(response, f'data-legacy-adoption="{meeting.pk}"')
        self.assertEqual(self.post(meeting).status_code, 404)
        with self.assertRaises(ValidationError):
            CoverageAdoptionService.preview(actor=self.actor, meeting=meeting)
        saved.refresh_from_db()
        self.assertEqual((saved.faculty_user, saved.revision), (self.replacement, 1))
        self.assertEqual(list(saved.history.values()), history)
        # Another synthetic meeting with explicit substitution is also protected.
        other = MeetingService.generate(actor=self.actor, schedule_slot=meeting.schedule_slot,
            meeting_date=date(2026, 1, 12), offerings=[])
        type(other).objects.filter(pk=other.pk).update(faculty_user=None, unresolved_coverage=True)
        other.refresh_from_db()
        SubstitutionService.assign(actor=self.actor, meeting=other, substitute_faculty=self.replacement, reason='')
        self.assertEqual(self.post(other).status_code, 404)
        self.assertFalse(MeetingCoverageAdoption.objects.exists())

    def test_publication_closure_and_unrelated_pending_review_remain_protected(self):
        meeting = self.ready()
        review = existing.MeetingReconciliation.objects.create(meeting=meeting, source_type='SCHEDULE',
            source_reference='synthetic-unrelated', status='PENDING', reason='Synthetic source review', detected_by=self.actor,
            before_snapshot=_meeting_snapshot(meeting), proposed_snapshot={'synthetic_schedule_review': True})
        self.assertContains(self.post(meeting), 'Resolve other source/history changes')
        self.assertFalse(MeetingCoverageAdoption.objects.exists())
        ReconciliationService.resolve(actor=self.actor, reconciliation=review, decision='KEEP_SNAPSHOT', reason='')
        # Test legacy protected evidence directly, only inside this rolled-back test DB.
        closure = AttendanceClosureDecision.objects.create(meeting=meeting, status='CLOSED', kind='HOLIDAY',
            pay_basis='REGULAR', faculty_user=self.faculty, decided_by=self.actor, reason='Synthetic legacy closure')
        self.assertContains(self.post(meeting), 'Recorded or published attendance')
        self.assertFalse(MeetingCoverageAdoption.objects.exists())
        publication = AttendanceCutoffPublication.objects.create(tenant=self.tenant, campus=self.campus,
            academic_year=self.academic_year, term=self.term, start_date=meeting.meeting_date,
            end_date=meeting.meeting_date, published_by=self.actor, review_fingerprint='synthetic',
            published_at=existing.timezone.now(), submission_key='synthetic-protected-publication')
        entry = AttendanceCutoffPublicationEntry.objects.create(publication=publication, meeting=meeting,
            occurrence_key='synthetic-protected', meeting_date=meeting.meeting_date, starts_at=meeting.starts_at,
            ends_at=meeting.ends_at, scheduled_minutes=60, status='PRESENT', faculty_user=self.faculty)
        # Remove the disposable fixture closure to exercise publication alone (not the recovery action).
        closure.delete()
        self.assertContains(self.post(meeting), 'Recorded or published attendance')
        response = self.client.get(reverse('faculty_attendance:reconciliation'))
        self.assertNotContains(response, f'data-legacy-adoption="{meeting.pk}"')
        entry.refresh_from_db()
        self.assertEqual(entry.faculty_user, self.faculty)
        self.assertFalse(MeetingCoverageAdoption.objects.exists())

    def test_missing_primary_link_has_actionable_error_without_guessing_a_candidate(self):
        meeting = self.ready()
        meeting.offering_links.update(is_primary=False)
        response = self.client.get(reverse('faculty_attendance:reconciliation'))
        self.assertContains(response, 'exactly one primary offering link')
        self.assertNotContains(response, 'Verified candidate:')
        self.assertNotContains(response, f'data-legacy-adoption="{meeting.pk}"')
        self.assertContains(self.post(meeting), 'exactly one primary offering link')
        self.assertFalse(MeetingCoverageAdoption.objects.exists())

    def sync_marked_meeting(self):
        meeting = self.ready()
        meeting.faculty_snapshot = {'college_academic_sync': True, 'note': 'Synthetic protected sync evidence'}
        meeting.save(update_fields=['faculty_snapshot'])
        return meeting

    def test_unresolved_sync_marker_get_has_diagnostics_without_verified_candidate_or_form(self):
        meeting = self.sync_marked_meeting()
        before = deepcopy(_meeting_snapshot(meeting))
        audits = existing.AuditLog.objects.count()
        response = self.client.get(reverse('faculty_attendance:reconciliation'))
        self.assertContains(response, 'College synchronization evidence')
        self.assertContains(response, 'Effective from')
        self.assertContains(response, 'Candidate not verified')
        self.assertNotContains(response, 'Verified candidate:')
        self.assertNotContains(response, f'data-legacy-adoption="{meeting.pk}"')
        self.assertNotContains(response, 'Adopt verified dated faculty')
        self.assertFalse(MeetingCoverageAdoption.objects.exists())
        meeting.refresh_from_db()
        self.assertEqual(_meeting_snapshot(meeting), before)
        self.assertEqual(existing.AuditLog.objects.count(), audits)

    def test_unresolved_sync_marker_post_rejected_without_mutating_frozen_evidence_or_history(self):
        meeting = self.sync_marked_meeting()
        round_ = CheckingRoundService.create(actor=self.actor, meetings=[meeting], checking_date=meeting.meeting_date)
        before = deepcopy(_meeting_snapshot(meeting))
        frozen = deepcopy(round_.manifest_rows.get().meeting_snapshot)
        reviews = list(meeting.reconciliations.values())
        coverages = list(FacultyCoverage.objects.values())
        audits = existing.AuditLog.objects.count()
        response = self.post(meeting)
        self.assertContains(response, 'Recovery was not saved.')
        self.assertContains(response, 'College synchronization evidence')
        self.assertNotContains(response, f'data-legacy-adoption="{meeting.pk}"')
        self.assertFalse(MeetingCoverageAdoption.objects.exists())
        self.assertFalse(AttendanceResult.objects.exists())
        meeting.refresh_from_db()
        self.assertEqual(_meeting_snapshot(meeting), before)
        self.assertEqual(round_.manifest_rows.get().meeting_snapshot, frozen)
        self.assertEqual(list(meeting.reconciliations.values()), reviews)
        self.assertEqual(list(FacultyCoverage.objects.values()), coverages)
        self.assertEqual(existing.AuditLog.objects.count(), audits)
        with self.assertRaises(ValidationError):
            require_confirmable_meeting_faculty(meeting)

    def test_truthy_unresolved_sync_marker_rejects_both_service_entry_points(self):
        meeting = self.sync_marked_meeting()
        for marker in (True, 'synthetic-sync', 1):
            with self.subTest(marker=marker):
                meeting.faculty_snapshot['college_academic_sync'] = marker
                meeting.save(update_fields=['faculty_snapshot'])
                before = deepcopy(_meeting_snapshot(meeting))
                audits = existing.AuditLog.objects.count()
                for operation in (CoverageAdoptionService.preview, CoverageAdoptionService.adopt):
                    with self.assertRaisesMessage(ValidationError, 'College synchronization evidence'):
                        operation(actor=self.actor, meeting=meeting)
                self.assertFalse(MeetingCoverageAdoption.objects.exists())
                meeting.refresh_from_db()
                self.assertEqual(_meeting_snapshot(meeting), before)
                self.assertEqual(existing.AuditLog.objects.count(), audits)


class ChecklistPrintExportTests(TestCase):
    permission_codes = CoverageInitializationTests.permission_codes
    setUpTestData = classmethod(CoverageInitializationTests.setUpTestData.__func__)
    setUp = CoverageInitializationTests.setUp
    aware = CoverageInitializationTests.aware
    schedule = CoverageInitializationTests.schedule
    coverage = CoverageInitializationTests.coverage
    meeting = CoverageInitializationTests.meeting
    assignment = CoverageInitializationTests.assignment
    unresolved = CoverageInitializationTests.unresolved

    def query(self, **kwargs):
        return dict(academic_year=self.academic_year.pk, term=self.term.pk, month='2026-01', day_group='MW', **kwargs)

    def get(self, route, **kwargs):
        self.client.force_login(self.actor)
        return self.client.get(reverse('faculty_attendance:' + route), self.query(**kwargs))

    def test_print_excludes_genuinely_unassigned_but_keeps_assigned_unresolved(self):
        self.assignment()
        response = self.get('monthly_print')
        self.assertEqual(response.status_code, 200)
        self.assertEqual({r.offering_id for r in response.context['rows']}, {self.offering.pk})
        self.assertContains(response, 'Historical faculty not verified')
        self.assertEqual(len(self.get('checklist').context['rows']), 2)
        self.assertFalse(FacultyCoverage.objects.exists())

    def test_dated_coverage_without_current_assignment_is_printable(self):
        self.coverage()
        self.assertEqual({r.offering_id for r in self.get('monthly_print').context['rows']}, {self.offering.pk})

    def test_explicit_substituted_combined_meeting_is_printed_and_exported_once(self):
        meeting = self.unresolved(combined=True)
        SubstitutionService.assign(actor=self.actor, meeting=meeting, substitute_faculty=self.faculty, reason='')
        response = self.get('monthly_print')
        rows = response.context['rows']
        self.assertEqual(len(rows), 1)
        self.assertTrue(rows[0].is_combined)
        self.assertEqual(rows[0].applicable_dates, frozenset([date(2026, 1, 5)]))
        exported = self.get('monthly_export')
        workbook = load_workbook(BytesIO(exported.content))
        self.assertEqual(workbook.active['B6'].value, 'S1\nS2')
        self.assertIsNone(workbook.active['F6'].value)
        self.assertEqual(workbook.active['G6'].value, 'N/A')

    def test_unassigned_combined_meeting_does_not_bypass_print_exclusion(self):
        self.unresolved(combined=True)
        self.assertEqual(self.get('monthly_print').context['rows'], [])

    def test_paper_choices_readable_defaults_validation_and_filter_retention(self):
        self.assignment()
        default = self.get('monthly_print')
        self.assertEqual(default.context['print_settings']['text_size'], 11)
        self.assertContains(default, 'font-size:11pt')
        for paper in ('A4', 'Letter', 'Legal', 'Long Bond'):
            for orientation in ('portrait', 'landscape'):
                response = self.get('monthly_print', paper=paper, orientation=orientation, text_size='14')
                self.assertEqual(response.status_code, 200)
                settings = response.context['print_settings']
                self.assertEqual((settings['paper'], settings['orientation'], settings['text_size']), (paper, orientation, 14))
                self.assertTrue(all(len(batch) <= settings['dates_per_sheet'] for batch in response.context['date_batches']))
                self.assertContains(response, 'font-size:14pt')
                self.assertIn('month=2026-01', response.context['filter_query'])
        invalid = self.get('checklist', paper='Unknown', text_size='6')
        self.assertTrue(invalid.context['monthly_form'].errors)
        self.assertEqual(self.get('monthly_export', paper='Unknown').status_code, 400)

    def test_xlsx_saved_order_literal_text_blank_cells_dates_and_print_settings(self):
        self.assignment()
        self.assignment(offering=self.combined_offering)
        # Formula-like source labels remain literal strings in the workbook.
        self.course.title = '=HYPERLINK("https://invalid.example")'
        self.course.save()
        response = self.get('checklist')
        tokens = [r.token for r in reversed(response.context['rows'])]
        MonthlyArrangementService.save(actor=self.actor, tenant_id=self.tenant.pk, campus_id=self.campus.pk,
            academic_year_id=self.academic_year.pk, term_id=self.term.pk, day_group='MW', tokens=tokens, expected_revision=None)
        printable = self.get('monthly_print', paper='Long Bond', orientation='portrait', text_size='12')
        exported = self.get('monthly_export', paper='Long Bond', orientation='portrait', text_size='12')
        self.assertEqual(exported.status_code, 200)
        workbook = load_workbook(BytesIO(exported.content))
        sheet = workbook.active
        self.assertEqual([sheet['B6'].value, sheet['B7'].value], [r.offering.section.code for r in printable.context['rows']])
        self.assertEqual(sheet['A6'].data_type, 's')
        self.assertIn('=HYPERLINK', sheet['A6'].value)
        self.assertIsNone(sheet['F6'].value)
        self.assertEqual(sheet['G6'].value, 'N/A')
        self.assertEqual(sheet['F4'].value.date(), date(2026, 1, 5))
        self.assertEqual((sheet.page_setup.paperSize, sheet.page_setup.orientation, sheet['B6'].font.sz), (14, 'portrait', 12))
        self.assertEqual(sheet.page_setup.scale, 100)
        self.assertEqual(sheet.print_title_rows, '$1:$4')
        self.assertEqual(len(workbook.worksheets), len(printable.context['date_batches']))

    def test_export_direct_deny_and_post_do_not_bypass_print_authority(self):
        self.assignment()
        self.client.force_login(self.actor)
        url = reverse('faculty_attendance:monthly_export')
        self.assertEqual(self.client.post(url, self.query()).status_code, 405)
        UserPermission.objects.create(user=self.actor, permission=Permission.objects.get(code=PRINT_PERMISSION),
            grant_type='DENY', tenant=None, campus=None)
        self.assertEqual(self.client.get(url, self.query()).status_code, 403)
